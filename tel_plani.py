# -*- coding: utf-8 -*-
"""
tel_plani.py — TK2 MEKANİZMA TELLERİ PLANI (kullanıcı 2026-10-09).

İSTEK: "TK2'de mekanizmalarda kullandığımız teller için bir plan hazırlayıp fasondan
sorumlu kişiden talep edeceğim … 2 haftalık, 4 haftalık." Modül, TEDARİKÇİLERLE
ilgilenen arkadaşa tel ihtiyacını göndermek için: kim üretiyor, hangi hafta kaç adet.

HESAP — TK2 MEKANİZMA İHTİYACINDAN TÜRETİLİR
--------------------------------------------
  1. Mekanizmalar = Forge referans listesi, bölüm=montaj, TK2 (93.* tel kodları hariç —
     TK2'de tel üretilmez, satılmaz; montaj planıyla aynı küme).
  2. Mekanizmaların açık emirleri (XPRO90: 10 = OPR, 40/45/50 = açık launch, kalan
     = sipariş − teyit). Emir ERP'de stok düşülerek açıldığı için mekanizma stoğu
     ikinci kez düşülmez (montaj planındaki 'opr_net' kuralı).
  3. Ürün ağacı (BSPEF2); HAYALİ ara düğümler 3 seviyeye kadar açılır, 93.* kod
     hayali olsa da açılmaz (tel kendisi teslim edilen kalemdir). Ağaçta ilk
     karşılaşılan 93.* kod mekanizmanın TELİDİR — W'li 'montajsız tel grubu' da
     olabilir (planlamanın MONTAJ TEL sayfasındaki 93.TK.030W gibi).
  4. Tel ihtiyacı = Σ mekanizma emri kalanı × ağaçtaki birim miktar, emrin
     bitiş (Q0FPD) haftasına göre: Gecikmiş · bu hafta · +1 · +2 · +3.
  5. Arz = tel stoğu (montaj planının sayılan depoları, aynı netleşme kuralı)
     + yoldaki tel launch'ları (40/45/50 kalan). Arz TARİH SIRASIYLA tüketilir;
     karşılanmayan kısım o haftanın TALEBİDİR.

NEDEN ERP'NİN TEL EMİRLERİ DEĞİL: XPRO90'daki 93.* emirleri ERP'nin kendi MRP
sonucu — tel stoğunu düşüyor ve TK1 ürünlerinin ihtiyacını da içeriyor (ölçüm
2026-10-09, laptop kopyası: 93.01.1307/20 TK2 ihtiyacı 100, ERP tel emri 873).
Panelde bilgi olarak gösterilir.

KİM ÜRETİYOR — ANAVERİ (kullanıcı 2026-10-09: "evet Anaveri'den alalım, Pandora
bir tedarikçi")
--------------------------------------------------------------------------------
Planlamanın PLAN 26xxxx.xlsb → 'Anaveri' sayfası: CD ART. · A/P · TK-1/2 · Hat ·
Makine · Kapasite 1V/Adet · zor/kolay. 93.* satırlarında TK-1/2 = Pandora (tedarikçi)
· TK-1 (içeride, ağırlıkla PP hattı) · Satınalma. İlk sürümdeki 'Fason' sayfası
(Zümel / Kuzey / Erkunt-Başak) ESKİMİŞTİ: Zümel listesindeki 387 kodun 281'i
Anaveri'de Pandora. Anaveri ortak tablodur (anaveri.py — satış planı da kullanır).

KONTROL — "Pull telleri şu anda fasonda üretiliyor" (kullanıcı 2026-10-09):
  · Hat = Pull → TEDARİKÇİ teli sayılır (Anaveri TK-1 dese de; o zaman uyarı).
  · TK-1/2 = TK-1 ama Hat = Pandora (80 satır, ör. 93.TK.1186) → çelişki uyarısı.
  · Anaveri'de yok → uyarı (W'li kod W'siz kodun satırını alır).
  · Tedarikçi teli son 30 günde TK1'de (tel/montaj) KAPAMA ya da SON MONTAJ
    görmüşse → uyarı ("içeride mi üretiliyor?"). Yalnız kesim/soyma ön işlemse
    bilgi notu (laptop verisi, Eylül: Pull teli sayılan 156 kodun TK1'de kaydı var,
    çoğu kesim/soyma; 93.00.1377 gibi bazılarında kapama da var).

AS400'e YALNIZ SELECT gider. Ölçüm JSON olarak saklanır; görünüm (2/4 hafta,
tedarikçi süzgeci) her istekte bu JSON'dan hesaplanır — AS400'e tekrar gidilmez.
"""
import json
import os
import re
import threading
import time
from datetime import date, datetime, timedelta

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
import anaveri as AV                           # planlamanın Anaveri'si — ortak tablo
VARSAYILAN_PLAN_KLASORU = AV.VARSAYILAN_PLAN_KLASORU
TEL_ONEK = ('93.',)
ACIK_DURUMLAR = ('10', '40', '45', '50')
SAKLANAN_OLCUM = 30
ICERIDE = ('TK-1', 'TK-2', 'TK1', 'TK2')
# Panelde elle seçilebilen değerler (serbest metin de girilebilir). 'TK1' = içeride.
SECENEKLER = ('Pandora', 'Satınalma', 'TK1')
# TK1'de bu adımlar görülmüşse tel İÇERİDE üretilmiş demektir (kesim/soyma ön işlem)
TAM_URETIM_ADIMLARI = ('KAPAMA', 'SON MONTAJ')
_GUVENLI_KOD = re.compile(r'^[A-Za-z0-9./\- ]{3,40}$')


# ── TABLOLAR ─────────────────────────────────────────────────────────────────
def tablolari_kur(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS tel_plani_olcum (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, kullanici TEXT DEFAULT '',
        mekanizma INTEGER DEFAULT 0, emirli INTEGER DEFAULT 0, tel INTEGER DEFAULT 0,
        sure_sn REAL, veri TEXT NOT NULL)""")
    # Elle seçim (Anaveri'yi ezer). Eski sürümün 'Fason' sayfası tohumu (kaynak='excel')
    # ESKİMİŞ bilgi — silinir; elle girilenler kalır.
    conn.execute("""CREATE TABLE IF NOT EXISTS tel_fason (
        kod TEXT PRIMARY KEY, fasoncu TEXT DEFAULT '', adaylar TEXT DEFAULT '',
        kaynak TEXT DEFAULT '', guncelleyen TEXT DEFAULT '', guncellendi TEXT DEFAULT '')""")
    conn.execute("DELETE FROM tel_fason WHERE kaynak <> 'elle' OR COALESCE(fasoncu,'') = ''")
    AV.tablolari_kur(conn)
    conn.execute("""CREATE TABLE IF NOT EXISTS tel_plani_talep (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, kullanici TEXT DEFAULT '',
        fasoncu TEXT DEFAULT '', ufuk INTEGER DEFAULT 4, launch_dahil INTEGER DEFAULT 1,
        olcum_ts TEXT DEFAULT '', notu TEXT DEFAULT '', kod_sayisi INTEGER DEFAULT 0,
        toplam REAL DEFAULT 0, haftalar TEXT DEFAULT '[]', satir TEXT NOT NULL)""")
    conn.commit()


# ── ANAVERİ (ortak tablo: anaveri.py) ──────────────────────────────────────
def anaveri_dosyadan(yol):
    return AV.dosyadan(yol)


def en_yeni_plan(klasor=VARSAYILAN_PLAN_KLASORU):
    return AV.en_yeni_plan(klasor)


def anaveri_yaz(conn, satirlar, dosya, kullanici, commit=True):
    return AV.yaz(conn, satirlar, dosya, kullanici, commit)


def anaveri_haritasi(conn):
    """Yalnız teller (93.*) — tedarikçi / Pull kontrolü için."""
    return AV.haritasi(conn, '93.')


def anaveri_durumu(conn):
    return AV.durumu(conn)


def elle_haritasi(conn):
    tablolari_kur(conn)
    return {r[0]: r[1] for r in conn.execute("SELECT kod, fasoncu FROM tel_fason WHERE kaynak='elle' "
                                             "AND COALESCE(fasoncu,'') <> ''")}


def tk1_tel_uretimi(conn, gun=30):
    """Son N günde TK1'de (tel + montaj) kaydedilen 93.* üretim: {kök kod: {adım: adet}}.
    Tel kayıtları adım ekiyle durur ('93.TK.464 KAPAMA') — kök ilk sözcüktür."""
    sinir = (date.today() - timedelta(days=gun)).isoformat()
    out = {}
    for ref, n in conn.execute(
            "SELECT u.referans_kodu, SUM(COALESCE(u.ok_adet,0)) FROM uretim_kayitlari u "
            "JOIN vardiyalar v ON v.id=u.vardiya_id WHERE COALESCE(v.lokasyon,'TK2')='TK1' "
            "AND v.bolum IN ('tel','montaj') AND v.tarih >= ? AND u.referans_kodu LIKE '93.%' "
            "GROUP BY u.referans_kodu", (sinir,)):
        p = str(ref or '').strip().split(None, 1)
        if not p or not n:
            continue
        adim = (p[1] if len(p) > 1 else 'ADIMSIZ').upper()
        d = out.setdefault(p[0].upper(), {})
        d[adim] = d.get(adim, 0) + n
    return out


def kaynak_bul(kod, anaveri, elle, tk1_uretim):
    """Tel kimden gelir? → {'tedarikci', 'sinif' (tedarikci|tk1|belirsiz), 'ana', 'uyarilar', 'notlar', 'elle'}"""
    uyar, notlar = [], []
    a, kok = anaveri.get(kod), kod
    if not a and kod.endswith('W') and anaveri.get(kod[:-1]):
        a, kok = anaveri.get(kod[:-1]), kod[:-1]
        notlar.append(f"Anaveri satırı W'siz koddan ({kod[:-1]})")
    tk, hat = (a or {}).get('tk', ''), (a or {}).get('hat', '')
    pull = hat.upper() == 'PULL'
    if not a:
        sinif, ted = 'belirsiz', ''
        uyar.append("Anaveri'de yok — kimin ürettiği bilinmiyor")
    elif pull:
        sinif = 'tedarikci'
        ted = '' if tk.upper() in ICERIDE or not tk else tk
        if tk.upper() in ICERIDE:
            uyar.append(f"Pull teli (şu an fasonda üretiliyor) ama Anaveri TK-1/2 '{tk}' diyor — tedarikçi seçin")
    elif tk.upper() in ICERIDE:
        sinif, ted = 'tk1', 'TK1'
        if hat.upper() == 'PANDORA':
            uyar.append(f"Anaveri çelişkili: TK-1/2 '{tk}' ama Hat 'Pandora' — içeride mi, Pandora'da mı?")
    elif tk:
        sinif, ted = 'tedarikci', tk
    else:
        sinif, ted = 'belirsiz', ''
        uyar.append("Anaveri'de TK-1/2 boş")
    e = elle.get(kod)
    if e:
        ted = e
        sinif = 'tk1' if e.upper() in ICERIDE else 'tedarikci'
    # TK1'de üretim izi: tedarikçiden gelmesi beklenen tel içeride yapılıyor mu?
    iz = tk1_uretim.get(kod) or (tk1_uretim.get(kok) if kok != kod else None) or {}
    if iz and sinif == 'tedarikci':
        tam = {k: v for k, v in iz.items() if any(x in k for x in TAM_URETIM_ADIMLARI)}
        ozet = ', '.join(f'{k.title()} {v:g}' for k, v in sorted(iz.items(), key=lambda kv: -kv[1])[:4])
        if tam:
            uyar.append(f"Son 30 günde TK1'de üretilmiş ({ozet}) — tedarikçiden mi, içeriden mi?")
        else:
            notlar.append(f"TK1'de ön işlem yapılıyor ({ozet})")
    return {'tedarikci': ted, 'sinif': sinif, 'ana': a, 'uyarilar': uyar, 'notlar': notlar, 'elle': bool(e)}


# ── AS400 ────────────────────────────────────────────────────────────────────
def urun_bilgisi(cn, kodlar):
    """{kod: {'aciklama', 'prov'}} — BARTF0."""
    out = {}
    kodlar = [k for k in kodlar if _GUVENLI_KOD.match(k)]
    cu = cn.cursor()
    for i in range(0, len(kodlar), 50):
        grup = kodlar[i:i + 50]
        cu.execute("SELECT TRIM(A0ARTI), TRIM(A0ARDS), TRIM(A0PROV) FROM TKC0301F.BARTF0 "
                   f"WHERE A0ARTI IN ({','.join('?' * len(grup))})", grup)
        for kod, ad, prov in cu.fetchall():
            out[kod] = {'aciklama': ad or '', 'prov': prov or ''}
    return out


def eldeki_stok(depolar, sayilan, eksi_dusulen, ayri=('CF2',)):
    """Montaj planındaki kural (kaynak_plan_kontrol.hesapla): CF2 ayrı sayılır; diğer
    sayılan depolar 01D ile netleşir; REP/02 eksisi serbest stoktan düşülür."""
    tuketilmis = sum(min(0.0, depolar.get(d, 0)) for d in eksi_dusulen if d not in sayilan)
    return (max(0.0, sum(depolar.get(d, 0) for d in sayilan if d not in ayri) + tuketilmis)
            + sum(max(0.0, depolar.get(d, 0)) for d in sayilan if d in ayri))


def olc(cn, kp, mekanizmalar, sayilan, eksi_dusulen):
    """AS400'den okur, saklanacak ölçümü (dict) döner. cn: açık ERP bağlantısı,
    kp: kaynak_plan_kontrol modülü (ağaç / stok / emir sorguları orada)."""
    mekanizmalar = sorted({str(k).strip().upper() for k in mekanizmalar
                           if k and not str(k).strip().upper().startswith(TEL_ONEK)})
    m_opr = kp.opr_ihtiyaclari(cn, mekanizmalar, 35)
    agac, iz = kp.urun_agaci_hayali(cn, mekanizmalar, azami_seviye=3, acilmayan_onekler=TEL_ONEK)
    teller = {}
    for mk, lst in agac.items():
        for alt, birim, um in lst:
            if not alt.startswith(TEL_ONEK):
                continue
            z = (iz.get(mk) or {}).get(alt) or {}
            teller.setdefault(alt, []).append({'kod': mk, 'birim': float(birim or 0), 'um': um or '',
                                               'seviye': z.get('seviye', 1), 'yol': z.get('yol', '')})
    kodlar = sorted(teller)
    stok = kp.stoklar(cn, kodlar) if kodlar else {}
    t_opr = kp.opr_ihtiyaclari(cn, kodlar, 35) if kodlar else {}
    bilgi = urun_bilgisi(cn, kodlar) if kodlar else {}
    tel_veri = {}
    for k in kodlar:
        dp = {d: v for d, v in (stok.get(k) or {}).items() if v}
        o = t_opr.get(k) or {}
        tel_veri[k] = {
            'aciklama': (bilgi.get(k) or {}).get('aciklama', ''),
            'prov': (bilgi.get(k) or {}).get('prov', ''),
            'depolar': dp,
            'stok': round(eldeki_stok(dp, sayilan, eksi_dusulen), 3),
            'eksi': any(dp.get(d, 0) < 0 for d in sayilan),
            'yolda': round(float(o.get('launch_adet') or 0), 3),          # açık tel launch (40/45/50)
            'tel_opr': round(sum(x['kalan'] for x in o.get('satirlar') or [] if x['durum'] == '10'), 3),
            'tel_emirler': [{'t': x['tarih'], 'k': x['kalan'], 'd': x['durum'], 'no': x['opr']}
                            for x in (o.get('satirlar') or [])][:30],
            'ustler': sorted(teller[k], key=lambda x: x['kod']),
        }
    emirli = {u['kod'] for lst in teller.values() for u in lst if u['kod'] in m_opr}
    return {
        'ts': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'bugun': date.today().isoformat(),
        'sayilan_depolar': list(sayilan), 'eksi_dusulen': list(eksi_dusulen),
        'mekanizma_sayisi': len(mekanizmalar), 'emirli_mekanizma': len(m_opr),
        'telli_mekanizma': len({u['kod'] for lst in teller.values() for u in lst}),
        'teller': tel_veri,
        # Yalnız teli olan mekanizmaların emirleri saklanır
        'mekanizmalar': {mk: [{'t': x['tarih'], 'k': x['kalan'], 'd': x['durum'], 'no': x['opr']}
                              for x in (m_opr.get(mk) or {}).get('satirlar') or []]
                         for mk in sorted(emirli)},
    }


def kaydet(conn, veri, kullanici='otomatik', sure_sn=None):
    tablolari_kur(conn)
    cur = conn.execute(
        "INSERT INTO tel_plani_olcum (ts, kullanici, mekanizma, emirli, tel, sure_sn, veri) VALUES (?,?,?,?,?,?,?)",
        (veri['ts'], kullanici, veri['mekanizma_sayisi'], veri['emirli_mekanizma'], len(veri['teller']),
         sure_sn, json.dumps(veri, ensure_ascii=False)))
    conn.execute("DELETE FROM tel_plani_olcum WHERE id NOT IN (SELECT id FROM tel_plani_olcum "
                 "ORDER BY id DESC LIMIT ?)", (SAKLANAN_OLCUM,))
    conn.commit()
    return cur.lastrowid


def son_olcum(conn):
    tablolari_kur(conn)
    r = conn.execute("SELECT id, ts, kullanici, sure_sn, veri FROM tel_plani_olcum ORDER BY id DESC LIMIT 1").fetchone()
    if not r:
        return None
    v = json.loads(r[4])
    v['_id'], v['_kullanici'], v['_sure_sn'] = r[0], r[1], r[3]
    return v


# ── ARKA PLAN KOŞUSU (panel düğmesi + 07:15/13:15) ─────────────────────────
_DENEME = {'calisiyor': False, 'basladi': None, 'bitti': None, 'kim': '', 'hata': '', 'sonuc': None,
           'sure_sn': None}
_KILIT = threading.Lock()


def deneme_durumu():
    return dict(_DENEME)


def tazele(conn_ac, kullanici, mekanizma_oku, erp_ac, kp, sayilan, eksi_dusulen):
    """Tam tur — senkron. conn_ac(): yeni SQLite bağlantısı; mekanizma_oku(conn) →
    kod listesi; erp_ac() → AS400 bağlantısı."""
    if not _KILIT.acquire(blocking=False):
        raise RuntimeError('Tel planı zaten tazeleniyor')
    t0 = time.time()
    _DENEME.update(calisiyor=True, basladi=datetime.now().strftime('%Y-%m-%d %H:%M:%S'), bitti=None,
                   kim=kullanici, hata='', sonuc=None, sure_sn=None)
    conn = conn_ac()
    try:
        mek = mekanizma_oku(conn)
        if not mek:
            raise RuntimeError('Forge referans listesinde TK2 montaj mekanizması yok')
        cn = erp_ac()
        try:
            veri = olc(cn, kp, mek, sayilan, eksi_dusulen)
        finally:
            try:
                cn.close()
            except Exception:
                pass
        sure = round(time.time() - t0, 1)
        oid = kaydet(conn, veri, kullanici, sure)
        _DENEME['sonuc'] = {'id': oid, 'tel': len(veri['teller']), 'mekanizma': veri['mekanizma_sayisi'],
                            'emirli': veri['emirli_mekanizma']}
        return _DENEME['sonuc']
    except Exception as e:
        _DENEME['hata'] = str(e)
        raise
    finally:
        _DENEME.update(calisiyor=False, bitti=datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                       sure_sn=round(time.time() - t0, 1))
        try:
            conn.close()
        except Exception:
            pass
        _KILIT.release()


def tazele_arka_planda(*args):
    """Arka planda başlatır → True; zaten çalışıyorsa False (panel /durum'u sorar)."""
    if _DENEME['calisiyor'] or _KILIT.locked():
        return False
    _DENEME.update(calisiyor=True, basladi=datetime.now().strftime('%Y-%m-%d %H:%M:%S'), bitti=None,
                   hata='', sonuc=None)

    def _kos():
        try:
            tazele(*args)
        except Exception as e:
            print(f'[TEL-PLANI] tazeleme hatası: {e}')
    threading.Thread(target=_kos, daemon=True, name='tel-plani-tazele').start()
    return True


# ── GÖRÜNÜM (2 / 4 hafta) ──────────────────────────────────────────────────
def _pazartesi(d):
    return d - timedelta(days=d.weekday())


def haftalar(ufuk, bugun=None):
    """[{'anahtar', 'etiket', 'bas', 'bit'}] — 'GEC' + bu hafta dahil ufuk kadar ISO hafta."""
    bugun = bugun or date.today()
    p = _pazartesi(bugun)
    out = [{'anahtar': 'GEC', 'etiket': 'Gecikmiş', 'bas': '', 'bit': (bugun - timedelta(days=1)).isoformat()}]
    for i in range(ufuk):
        b = p + timedelta(days=7 * i)
        out.append({'anahtar': f'H{i}', 'etiket': f'{b.isocalendar()[1]}. hf',
                    'bas': max(b, bugun).isoformat(), 'bit': (b + timedelta(days=6)).isoformat()})
    return out


def _kova(t, bugun, pzt, ufuk):
    try:
        d = date.fromisoformat(t)
    except (TypeError, ValueError):
        return None
    if d < bugun:
        return 'GEC'
    i = (_pazartesi(d) - pzt).days // 7
    return f'H{i}' if 0 <= i < ufuk else None


def _suzgec_uyar(k, filtre):
    """'' tümü · '__tedarikci__' tedarikçiden gelenler · '__tk1__' içeride · '__yok__'
    tedarikçisi belirsiz · '__uyari__' kontrol uyarısı olanlar · aksi hâlde tedarikçi adı."""
    if not filtre:
        return True
    if filtre == '__tedarikci__':
        return k['sinif'] == 'tedarikci'
    if filtre == '__tk1__':
        return k['sinif'] == 'tk1'
    if filtre == '__yok__':
        return k['sinif'] != 'tk1' and not k['tedarikci']
    if filtre == '__uyari__':
        return bool(k['uyarilar'])
    return k['tedarikci'] == filtre


def son_talepler(conn, limit=200):
    """{kod: {'ts', 'adet', 'id'}} — her telin EN SON talebi."""
    tablolari_kur(conn)
    out = {}
    for tid, ts, satir in conn.execute("SELECT id, ts, satir FROM tel_plani_talep ORDER BY id DESC LIMIT ?", (limit,)):
        for s in json.loads(satir or '[]'):
            out.setdefault(s['kod'], {'ts': ts, 'adet': s.get('talep', 0), 'id': tid})
    return out


def gorunum(veri, anaveri, elle, tk1_uretim, ufuk=4, launch_dahil=True, fasoncu='', ara='', hepsi=False,
            talepler=None, bugun=None):
    """Ölçümden tablo satırları. ufuk: 2 ya da 4 hafta (bu hafta dahil).
    launch_dahil: launch'ı açık mekanizma emirleri de ihtiyaca girsin mi.
    fasoncu: tedarikçi süzgeci (bkz. _suzgec_uyar)."""
    bugun = bugun or date.today()
    pzt = _pazartesi(bugun)
    ufuk = 2 if int(ufuk or 4) <= 2 else 4
    hf = haftalar(ufuk, bugun)
    anahtarlar = [h['anahtar'] for h in hf]
    durumlar = ACIK_DURUMLAR if launch_dahil else ('10',)
    ara = re.sub(r'\s', '', str(ara or '')).upper()
    talepler = talepler or {}
    mek = veri.get('mekanizmalar') or {}
    satirlar = []
    for kod, t in (veri.get('teller') or {}).items():
        k = kaynak_bul(kod, anaveri, elle, tk1_uretim)
        if not _suzgec_uyar(k, fasoncu):
            continue
        if ara and ara not in kod.replace(' ', '').upper() and not any(
                ara in u['kod'].replace(' ', '').upper() for u in t['ustler']):
            continue
        ihtiyac = {a: 0.0 for a in anahtarlar}
        ustler = []
        for u in t['ustler']:
            u_iht, en_yakin = 0.0, ''
            for e in mek.get(u['kod']) or []:
                if e['d'] not in durumlar:
                    continue
                kv = _kova(e['t'], bugun, pzt, ufuk)
                if not kv:
                    continue
                q = float(e['k'] or 0) * u['birim']
                ihtiyac[kv] += q
                u_iht += q
                if e['t'] and (not en_yakin or e['t'] < en_yakin):
                    en_yakin = e['t']
            ustler.append({'kod': u['kod'], 'birim': u['birim'], 'yol': u.get('yol', ''),
                           'ihtiyac': round(u_iht, 2), 'en_yakin': en_yakin})
        toplam = sum(ihtiyac.values())
        arz = float(t['stok'] or 0) + float(t['yolda'] or 0)
        kalan, talep = arz, {}
        for a in anahtarlar:                       # arz TARİH SIRASIYLA tüketilir
            n = max(0.0, ihtiyac[a] - kalan)
            kalan = max(0.0, kalan - ihtiyac[a])
            talep[a] = round(n, 2)
        talep_top = round(sum(talep.values()), 2)
        if not hepsi and toplam <= 0:
            continue
        ustler.sort(key=lambda x: (-x['ihtiyac'], x['kod']))
        ana = k['ana'] or {}
        kap = ana.get('kapasite')
        satirlar.append({
            'kod': kod, 'aciklama': t.get('aciklama', ''), 'prov': t.get('prov', ''),
            'tedarikci': k['tedarikci'], 'sinif': k['sinif'], 'elle': k['elle'],
            'uyarilar': k['uyarilar'], 'notlar': k['notlar'],
            'hat': ana.get('hat', ''), 'makine': ana.get('makine', ''), 'zorluk': ana.get('zorluk', ''),
            'anaveri_tk': ana.get('tk', ''), 'kapasite': kap,
            'vardiya': round(talep_top / kap, 1) if kap and talep_top > 0 else None,
            'ihtiyac': {a: round(v, 2) for a, v in ihtiyac.items()}, 'ihtiyac_toplam': round(toplam, 2),
            'talep': talep, 'talep_toplam': talep_top,
            'stok': t['stok'], 'yolda': t['yolda'], 'tel_opr': t.get('tel_opr', 0),
            'depolar': t.get('depolar') or {}, 'eksi': t.get('eksi', False),
            'ustler': ustler, 'ust_sayisi': len(ustler),
            'son_talep': talepler.get(kod),
        })

    def ilk_talep(s):
        return next((i for i, a in enumerate(anahtarlar) if s['talep'][a] > 0), 99)
    satirlar.sort(key=lambda s: (0 if s['talep_toplam'] > 0 else 1, ilk_talep(s), -s['talep_toplam'], s['kod']))
    # Tedarikçi bazında toplam — tedarikçiyle ilgilenen arkadaşın ilk bakacağı yer
    ted = {}
    for s in satirlar:
        ad = s['tedarikci'] or ('TK1' if s['sinif'] == 'tk1' else '— belirsiz')
        d = ted.setdefault(ad, {'ad': ad, 'sinif': s['sinif'], 'tel': 0, 'talepli': 0, 'talep': 0.0,
                                'vardiya': 0.0, 'kapasitesiz': 0})
        d['tel'] += 1
        if s['talep_toplam'] > 0:
            d['talepli'] += 1
            d['talep'] += s['talep_toplam']
            if s['vardiya'] is None:
                d['kapasitesiz'] += 1
            else:
                d['vardiya'] += s['vardiya']
    return {
        'haftalar': hf, 'ufuk': ufuk, 'launch_dahil': bool(launch_dahil), 'satirlar': satirlar,
        'tedarikciler': sorted(({**d, 'talep': round(d['talep'], 2), 'vardiya': round(d['vardiya'], 1)}
                                for d in ted.values()), key=lambda d: -d['talep']),
        'ozet': {'tel': len(satirlar),
                 'talepli': sum(1 for s in satirlar if s['talep_toplam'] > 0),
                 'talep': round(sum(s['talep_toplam'] for s in satirlar), 2),
                 'gecikmis_talep': round(sum(s['talep']['GEC'] for s in satirlar), 2),
                 'ihtiyac': round(sum(s['ihtiyac_toplam'] for s in satirlar), 2),
                 'uyari': sum(1 for s in satirlar if s['uyarilar'])},
    }


# ── EXCEL ────────────────────────────────────────────────────────────────────
def excel(satirlar, hf, baslik, alt_baslik=''):
    """Talep listesi — tedarikçilerle ilgilenen arkadaşa gönderilecek biçim. bytes döner."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tel Talebi'
    ws.append([baslik])
    ws['A1'].font = Font(bold=True, size=13)
    if alt_baslik:
        ws.append([alt_baslik])
    ws.append([])
    hf_etiket = [h['etiket'] + (f" ({_tr(h['bas'])}–{_tr(h['bit'])})" if h['bas'] else '') for h in hf]
    bas = (['Tel kodu', 'Açıklama', 'Tedarikçi', 'Hat'] + hf_etiket
           + ['Toplam talep', 'Kapasite (1 vardiya)', 'Gerekli vardiya', 'Kullanıldığı mekanizmalar', 'Kontrol'])
    ws.append(bas)
    bas_satir = ws.max_row
    dolgu = PatternFill('solid', fgColor='E9E4F7')
    for i in range(1, len(bas) + 1):
        c = ws.cell(row=bas_satir, column=i)
        c.font, c.fill = Font(bold=True), dolgu
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for s in satirlar:
        ws.append([s['kod'], s.get('aciklama', ''), s.get('tedarikci') or '', s.get('hat', '')]
                  + [(s['talep'].get(h['anahtar']) or None) for h in hf]
                  + [s['talep_toplam'], s.get('kapasite'), s.get('vardiya'),
                     ', '.join(u['kod'] for u in s.get('ustler') or [] if u.get('ihtiyac')),
                     ' · '.join(s.get('uyarilar') or [])])
    ws.append([])
    ws.append(['TOPLAM', '', '', ''] + [sum(s['talep'].get(h['anahtar'], 0) for s in satirlar) or None for h in hf]
              + [sum(s['talep_toplam'] for s in satirlar)])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for i, gen in enumerate([16, 30, 13, 8] + [13] * len(hf) + [13, 12, 11, 40, 50], start=1):
        ws.column_dimensions[get_column_letter(i)].width = gen
    ws.freeze_panes = ws.cell(row=bas_satir + 1, column=2)
    # 2. sayfa: hesabın dökümü — talep neden bu kadar
    w2 = wb.create_sheet('Ayrıntı')
    w2.append(['Tel kodu', 'Mekanizma', 'Birim', 'Mekanizma ihtiyacından tel', 'En yakın emir',
               'Tel stoğu (sayılan)', 'Yolda (tel launch)', 'Toplam ihtiyaç', 'Toplam talep', 'Depolar'])
    for i in range(1, 11):
        w2.cell(row=1, column=i).font = Font(bold=True)
    for s in satirlar:
        dp = ' '.join(f'{d}:{v:g}' for d, v in (s.get('depolar') or {}).items())
        for j, u in enumerate(s.get('ustler') or []):
            w2.append([s['kod'], u['kod'], u['birim'], u['ihtiyac'], _tr(u.get('en_yakin', '')),
                       s['stok'] if j == 0 else None, s['yolda'] if j == 0 else None,
                       s['ihtiyac_toplam'] if j == 0 else None, s['talep_toplam'] if j == 0 else None,
                       dp if j == 0 else ''])
    for i, gen in enumerate([16, 18, 8, 14, 12, 14, 14, 13, 13, 40], start=1):
        w2.column_dimensions[get_column_letter(i)].width = gen
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _tr(iso):
    try:
        return datetime.strptime(iso, '%Y-%m-%d').strftime('%d.%m')
    except (TypeError, ValueError):
        return ''
