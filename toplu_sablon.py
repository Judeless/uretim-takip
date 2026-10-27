# -*- coding: utf-8 -*-
"""TOPLU REFERANS ŞABLONU (kullanıcı 2026-10-26, Aspanda): her bölüm AYRI SAYFA.

"Aspanda veri Excel'ini düzelt, her bölümü ayrı bir sayfa yap. O Excel'i Aspanda ile
paylaşacağım; toplu güncellemek isterlerse o şablon üzerinden ilerleyebilsinler."

Panel (Referanslar → Toplu Şablon / Toplu Yükle) ve tek seferlik dönüştürme betikleri
AYNI biçimi buradan yazar ve okur:
  · sayfa adı = bölümün görünen adı ('Metal Enjeksiyon', 'CNC İşleme' …)
  · 1. satır başlık; kod + açıklama + süre sütunları (bölüme göre: kaynakta kaynak /
    söktak, metal ve plastikte kalıp göz + BASKI süresi — parça başı süre ondan hesaplanır)
  · BOŞ hücre = dokunma; yeni kod = yeni referans; hiçbir satır SİLİNMEZ."""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GOZLU = ('metal', 'plastik')
SUTUNLAR = {        # anahtar → (başlık, eş adlar)
    'kod': ('Referans Kodu', ('referans kodu', 'referans', 'kod', 'kodu', 'ürün kodu', 'urun kodu',
                              'parça kodu', 'parca kodu', 'group/ext. code')),
    'aciklama': ('Açıklama', ('açıklama', 'aciklama', 'description', 'tanım', 'tanim')),
    'goz': ('Kalıp Göz', ('kalıp göz', 'kalip goz', 'göz', 'goz', 'göz adedi', 'goz adedi', '# of cavities')),
    'baski': ('Baskı Süresi (sn)', ('baskı süresi (sn)', 'baski suresi (sn)', 'baskı süresi', 'baski suresi',
                                    'çevrim süresi (baskı)', 'cevrim suresi (baski)')),
    'ks': ('Kaynak Süresi (sn)', ('kaynak süresi (sn)', 'kaynak suresi (sn)', 'kaynak süresi', 'kaynak suresi')),
    'ss': ('Söktak Süresi (sn)', ('söktak süresi (sn)', 'soktak süresi (sn)', 'soktak suresi (sn)',
                                  'söktak süresi', 'soktak suresi')),
    'ct': ('Parça Başı Süre (sn)', ('parça başı süre (sn)', 'parca basi sure (sn)', 'parça başı süre',
                                    'cycle time (sn)', 'cycle time', 'süre (sn)', 'sure (sn)', 'süre', 'sure')),
}
SUST = 100000.0
_TR = str.maketrans('çÇğĞıİöÖşŞüÜ', 'cCgGiIoOsSuU')


def katla(s):
    return ' '.join(str(s or '').translate(_TR).lower().split())


def kod_norm(s):
    return ''.join(str(s or '').split()).upper()


def sutunlar(bolum):
    s = ['kod', 'aciklama']
    if bolum in GOZLU:
        s += ['goz', 'baski']
    if bolum == 'kaynak':
        s += ['ks', 'ss']
    return s + ['ct']


def _sayi(v):
    """Hücre → (değer, hata). Boş = (None, None) 'dokunma'."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    try:
        f = float(str(v).strip().replace(',', '.'))
    except ValueError:
        return None, f'sayı değil: {v!r}'
    if f < 0 or f > SUST:
        return None, f'geçersiz değer: {f:g}'
    return f, None


def yaz(bolum_satirlari, adlar, baslik_notu=''):
    """bolum_satirlari: {bolum: [{'kod', 'aciklama', 'ct', 'goz', 'ks', 'ss'}]} (sıra korunur).
    adlar: {bolum: görünen ad}. → xlsx bayt."""
    wb = Workbook()
    wb.remove(wb.active)
    bas_font, bas_dolgu = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F2A44')
    bos_font = Font(color='B45309', bold=True)
    toplam_bos = {}
    for bolum, liste in bolum_satirlari.items():
        ws = wb.create_sheet((adlar.get(bolum) or bolum)[:31])
        anahtar = sutunlar(bolum)
        ws.append([SUTUNLAR[a][0] for a in anahtar])
        for h in ws[1]:
            h.font, h.fill = bas_font, bas_dolgu
            h.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        bos = 0
        for r in liste:
            ct = float(r.get('ct') or 0)
            goz = int(r.get('goz') or 1)
            deger = {'kod': r['kod'], 'aciklama': r.get('aciklama') or None,
                     'ct': round(ct, 2) or None, 'goz': goz,
                     'baski': round(ct * goz, 1) or None,      # parça başı 2 ondalık → baskı 95,04 değil 95
                     'ks': round(float(r.get('ks') or 0), 2) or None,
                     'ss': round(float(r.get('ss') or 0), 2) or None}
            ws.append([deger[a] for a in anahtar])
            if not ct:
                bos += 1
                ws.cell(row=ws.max_row, column=1).font = bos_font
        toplam_bos[bolum] = (bos, len(liste))
        for i, a in enumerate(anahtar, start=1):
            gen = {'kod': 22, 'aciklama': 48}.get(a, 16)
            ws.column_dimensions[get_column_letter(i)].width = gen
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(anahtar))}{max(2, ws.max_row)}'
    yg = wb.create_sheet('Nasıl doldurulur')
    satirlar = [
        ['TOPLU REFERANS ŞABLONU — her bölüm ayrı sayfada' + (f' · {baslik_notu}' if baslik_notu else '')],
        [''],
        ['1) Sayfa adlarını ve başlık satırını DEĞİŞTİRMEYİN; her bölümün referansları kendi sayfasındadır.'],
        ['2) BOŞ hücre = dokunma. Yalnız değiştirmek istediğiniz hücreyi doldurun; dosyayı parça parça doldurup'],
        ['   birden çok kez yükleyebilirsiniz.'],
        ['3) Listede olmayan YENİ bir kod yazarsanız yeni referans olarak eklenir. Hiçbir referans SİLİNMEZ.'],
        ['4) Süreler PARÇA BAŞINA saniyedir. Metal / plastikte "Baskı Süresi"ni ya da "Kalıp Göz"ü değiştirirseniz'],
        ['   parça başı süre = baskı süresi ÷ kalıp göz olarak hesaplanır; yalnız "Parça Başı Süre"yi değiştirirseniz o esas alınır.'],
        ['5) Panel → Referanslar → "⬆️ Toplu Yükle" ile dosyayı seçin: önce ÖNİZLEME gösterilir —'],
        ['   eklenecek / değişecek referansları görüp onayladığınızda uygulanır.'],
        ['6) Turuncu kodların süresi henüz tanımlı değil.'],
        [''],
    ] + [[f'{adlar.get(b, b)}: {n} referans, süresi boş {bos}'] for b, (bos, n) in toplam_bos.items()]
    for s in satirlar:
        yg.append(s)
    yg.column_dimensions['A'].width = 110
    yg['A1'].font = Font(bold=True, size=12)
    bellek = io.BytesIO()
    wb.save(bellek)
    return bellek.getvalue()


def oku(dosya, adlar):
    """Doldurulmuş şablon → ({bolum: [satır]}, hatalı[], tanınmayan sayfalar[]).
    satır: {'satir', 'kod', 'aciklama', 'ct', 'goz', 'ks', 'ss'} (None = dokunma)."""
    wb = load_workbook(dosya, data_only=True, read_only=True)
    ad_bolum = {}
    for b, ad in adlar.items():
        ad_bolum[katla(ad)] = b
        ad_bolum[katla(b)] = b
    sonuc, hatali, tanimsiz = {}, [], []
    for ws in wb.worksheets:
        bolum = ad_bolum.get(katla(ws.title))
        if not bolum:
            if katla(ws.title) != katla('Nasıl doldurulur'):
                tanimsiz.append(ws.title)
            continue
        it = ws.iter_rows(values_only=True)
        try:
            bas = [katla(x) for x in next(it)]
        except StopIteration:
            continue
        kol = {}
        for a, (_b, esler) in SUTUNLAR.items():
            for i, h in enumerate(bas):
                if h in {katla(e) for e in esler}:
                    kol[a] = i
                    break
        if 'kod' not in kol:
            hatali.append({'sayfa': ws.title, 'satir': 1, 'kod': '', 'sebep': "'Referans Kodu' başlığı yok"})
            continue
        liste = sonuc.setdefault(bolum, [])
        for no, r in enumerate(it, start=2):
            if not r:
                continue
            al = (lambda a: r[kol[a]] if (a in kol and kol[a] < len(r)) else None)
            kod = str(al('kod') or '').strip()
            if isinstance(al('kod'), float) and al('kod').is_integer():
                kod = str(int(al('kod')))
            if not kod:
                continue
            satir = {'satir': no, 'kod': kod, 'aciklama': (str(al('aciklama')).strip() or None)
                     if al('aciklama') not in (None, '') else None}
            hata = None
            for a in ('ct', 'goz', 'baski', 'ks', 'ss'):
                satir[a], e = _sayi(al(a))
                hata = hata or e
            if hata:
                hatali.append({'sayfa': ws.title, 'satir': no, 'kod': kod, 'sebep': hata})
                continue
            liste.append(satir)
    wb.close()
    return sonuc, hatali, tanimsiz
