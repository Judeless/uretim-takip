# -*- coding: utf-8 -*-
"""
aylik_uretim.py — AYLIK ÜRETİM RAPORU: teyidi verilmiş üretim, Forge bölümlerine göre
(kullanıcı 2026-10-09: "aylık üretilen, üretim teyidi verilen referansları bölümlere göre
ayırıp Excel hâlinde raporlayacağız … Forge bölümleri olsun, performans da eklensin,
çalışma saatini görelim; yönetici ile paylaşalım").

PLANLAMANIN ESKİ YÖNTEMİ
  · Q:\\UretimPlanlama\\EMRE\\Yeni klasör\\Üretim\\Yıllık Üretim\\2025 yıllık Üretim.xlsx —
    AS400 10-05-03-15 raporu (aylık), RPR ve CFI hareketleri; depo 01 = TK1, 01D = TK2.
  · Q:\\UretimPlanlama\\Aylık Kapasite Sunum\\Kapasite Kullanım Oranı 2026.xlsx — ayın kod
    × adet listesi 'Database' (kod → üretim hattı + saatlik adet) ile eşlenir; Summary:
    adet · teorik süre (adet ÷ saatlik adet) · çalışma saati ('Çalışma Saati' sayfası) ·
    performans (hedef %90). Database'de olmayan kod özete GİRMİYORDU (Ağustos 2026:
    677 kodun 226'sı, adedin %45'i; pres/abkant hiç yoktu).

FORGE
  · Adet: AS400 BMMAF0 — MGCACD 'RPR' (launch teyidi) + 'CFI' (launch'sız üretim girişi),
    HAREKET tarihine göre ay (10-05-03-15 ile aynı). Tesis hareketin DEPOSUNDAN
    (01D → TK2, 01/02 → TK1); bilinmeyen depoda kodun Forge tesisi.
  · Bölüm: FORGE (referans_listesi; tel'de adım ekli satırlar kökten eşlenir). Kod bir
    tesiste birden çok bölümde tanımlıysa her bölümün işine sayılır (o bölüm o koda iş
    yaptı); tesis toplamı kodu BİR kez sayar. Forge'da yoksa planlamanın Anaveri'si
    (TK-1/2 · Hat), o da yoksa 'Tanımsız' listesi.
  · Teorik süre (kapasite modülüyle AYNI öncelik): kapasite Excel süresi (kapasite_sure)
    → Forge çevrim süresi (tel'de adımların toplamı) → ERP rota süresi → yok.
  · Çalışma saati: kapasite Excel'inin 'Çalışma Saati' sayfası (NçS + %90/%95/%100 mesai,
    bölüm adıyla). PP / Pull saatleri TK1 TEL bölümüne yazılır (PP telleri Forge'da tel).
  · Performans = teorik süre ÷ toplam çalışma saati (NçS + mesai). Planlamanın eski
    tablosu yalnız NçS'ye bölüyordu — o değer de ayrı sütunda.
  · 20.* / 21.* hammadde (metre, kg) adet toplamına girmez — ayrı listede.

AS400'e YALNIZ SELECT gider. Ayın okuması saklanır (aylik_uretim_cekim); sınıflama ve
süreler rapor anında güncel tanımlarla uygulanır.
"""
import json
import os
import re
import threading
import time
from datetime import date, datetime

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
HAREKET_TABLO = 'tkc0301F.BMMAF0'
CAUSALLAR = ('RPR', 'CFI')
DEPO_TESIS = {'01D': 'TK2', '01': 'TK1', '02': 'TK1'}
HAMMADDE_ONEK = ('20.', '21.')
HEDEF_PERFORMANS = 0.90
# Excel süresi KULLANILMAYAN bölümler (kapasite modülüyle aynı kural, kullanıcı 2026-09-17): metalde
# Excel'in saatlik adedi ÇEVRİM başına, Forge'unki parça başına (kalıp gözü) — Forge süresi önce gelir.
SURE_EXCEL_DISI = ('metal',)
VARSAYILAN_KAPASITE_KLASORU = r'Q:\UretimPlanlama\Aylık Kapasite Sunum'
# Kapasite Excel'i 'Çalışma Saati' BÖLÜM adı → Forge (tesis, bölüm)
CALISMA_BOLUM = {
    'WELDING': ('TK2', 'kaynak'), 'LASER CUTTING': ('TK2', 'lazer'), 'METAL INJECTION': ('TK2', 'metal'),
    'ASSEMBLY': ('TK2', 'montaj'), 'ABKANT': ('TK2', 'pres'),
    'PP': ('TK1', 'tel'), 'PULL': ('TK1', 'tel'),
    'IVECO': ('TK1', 'montaj'), 'JKP,JKW': ('TK1', 'montaj'), 'LF,LFP,LAG (LTK),LSB,PBL': ('TK1', 'montaj'),
    'PLASTIC INJECTION': ('TK1', 'plastik'),
}
# Anaveri (TK-1/2 · Hat) → Forge bölümü — yalnız kod Forge'da HİÇ yoksa
ANAVERI_BOLUM = {
    ('TK-2', 'MONTAJ'): ('TK2', 'montaj'), ('TK-2', 'TRAKTÖR'): ('TK2', 'montaj'), ('TK-2', 'PEDAL'): ('TK2', 'montaj'),
    ('TK-2', 'PROJE'): ('TK2', 'montaj'), ('TK-2', 'KAYNAK'): ('TK2', 'kaynak'), ('TK-2', 'LAZER KESIM'): ('TK2', 'lazer'),
    ('TK-2', 'LAZER KESİM'): ('TK2', 'lazer'), ('TK-2', 'METAL ENJEKSIYON'): ('TK2', 'metal'),
    ('TK-2', 'METAL ENJEKSİYON'): ('TK2', 'metal'), ('TK-2', 'SAC BÜKÜM'): ('TK2', 'pres'),
    ('TK-2', 'İŞLEME'): ('TK2', 'isleme'), ('TK-2', 'ISLEME'): ('TK2', 'isleme'),
    ('TK-1', 'PP'): ('TK1', 'tel'), ('TK-1', 'KESIM'): ('TK1', 'tel'), ('TK-1', 'KESİM'): ('TK1', 'tel'),
    ('TK-1', 'LF'): ('TK1', 'montaj'), ('TK-1', 'JOYSTICK'): ('TK1', 'montaj'), ('TK-1', 'IVECO'): ('TK1', 'montaj'),
    ('TK-1', 'PLASTIK ENJEKSIYON'): ('TK1', 'plastik'), ('TK-1', 'PLASTIK ENJEKSİYON'): ('TK1', 'plastik'),
    # Tabo hortumları (92.*) Forge'da TK1 Montaj
    ('TABO', 'B. HOSE'): ('TK1', 'montaj'),
}
# FASON / TEDARİKÇİ (kullanıcı 2026-10-09: "Pull telleri şu anda fasonda üretiliyor"): Anaveri'de
# Pandora / Pull tellerinin RPR-CFI'ı tedarikçinin işidir — TK1 Tel'e sayılsaydı teorik süresi TK1
# çalışanlarının saatine bölünür, performansı şişirirdi. Ayrı satır, performansa girmez.
FASON = 'fason'


def _fason_mi(a, kesin=True):
    """kesin=True: Pandora / Pull (Forge tanımını da ezer). False: Anaveri 'Fason' (yalnız
    Forge'da tanım yoksa — Forge'da tanımlıysa içeride üretiyoruz demektir)."""
    if not a:
        return False
    tk, hat = (a.get('tk') or '').upper(), (a.get('hat') or '').upper()
    if kesin:
        return tk == 'PANDORA' or hat == 'PULL'
    return tk == 'FASON' or hat == 'FASON'
BOLUM_AD = {'kaynak': 'Kaynak', 'montaj': 'Montaj', 'metal': 'Metal Enjeksiyon', 'lazer': 'Lazer Kesim',
            'pres': 'Pres / Abkant', 'isleme': 'İşleme', 'plastik': 'Plastik Enjeksiyon', 'tel': 'Tel Üretimi',
            'fason': 'Fason (tedarikçi)'}
BOLUM_SIRA = {'TK2': ('kaynak', 'lazer', 'pres', 'metal', 'isleme', 'montaj', 'fason'),
              'TK1': ('montaj', 'tel', 'plastik', 'fason')}
AY_AD = ('', 'Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim',
         'Kasım', 'Aralık')
_GUVENLI_KOD = re.compile(r'^[A-Za-z0-9./\- ]{3,40}$')


def tablolari_kur(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS aylik_uretim_cekim (
        id INTEGER PRIMARY KEY AUTOINCREMENT, yil INTEGER NOT NULL, ay INTEGER NOT NULL,
        ts TEXT NOT NULL, kullanici TEXT DEFAULT '', satir INTEGER DEFAULT 0, sure_sn REAL,
        veri TEXT NOT NULL)""")
    conn.execute("CREATE INDEX IF NOT EXISTS ix_auc_ay ON aylik_uretim_cekim(yil, ay)")
    conn.execute("""CREATE TABLE IF NOT EXISTS aylik_calisma (
        yil INTEGER NOT NULL, ay INTEGER NOT NULL, excel_bolum TEXT NOT NULL,
        nos REAL DEFAULT 0, mesai REAL DEFAULT 0, kisi INTEGER DEFAULT 0, isimsiz REAL DEFAULT 0,
        dosya TEXT DEFAULT '', kullanici TEXT DEFAULT '', guncellendi TEXT DEFAULT '',
        PRIMARY KEY (yil, ay, excel_bolum))""")
    # Kapasite Excel'inin 'Database' sayfası: kod → saatlik adet + üretim hattı (ayın değil, son yüklenen)
    conn.execute("""CREATE TABLE IF NOT EXISTS aylik_kapasite_db (
        kod TEXT PRIMARY KEY, saatlik REAL, hat TEXT DEFAULT '', makine TEXT DEFAULT '',
        dosya TEXT DEFAULT '', guncellendi TEXT DEFAULT '')""")
    conn.commit()


def _norm(k):
    return re.sub(r'\s+', '', str(k or '')).upper()


def _kok(k):
    return _norm(str(k or '').strip().split(' ')[0])


# ── AS400 ────────────────────────────────────────────────────────────────────
def as400_oku(cn, yil, ay):
    """Ayın RPR + CFI hareketleri, kod × depo × causal toplamı + ürün açıklaması."""
    cu = cn.cursor()
    cu.execute(f"SELECT TRIM(MGARCD), TRIM(MGMGCD), TRIM(MGCACD), SUM(MGQTA), COUNT(*) FROM {HAREKET_TABLO} "
               f"WHERE MGCACD IN ('RPR','CFI') AND MGDSSO=? AND MGDAAO=? AND MGDMMO=? "
               f"GROUP BY TRIM(MGARCD), TRIM(MGMGCD), TRIM(MGCACD)", (yil // 100, yil % 100, ay))
    hareket = [{'kod': r[0], 'depo': r[1] or '', 'causal': r[2], 'adet': float(r[3] or 0), 'satir': int(r[4] or 0)}
               for r in cu.fetchall() if r[0]]
    kodlar = sorted({h['kod'] for h in hareket if _GUVENLI_KOD.match(h['kod'])})
    aciklama = {}
    for i in range(0, len(kodlar), 50):
        grup = kodlar[i:i + 50]
        cu.execute("SELECT TRIM(A0ARTI), TRIM(A0ARDS) FROM TKC0301F.BARTF0 "
                   f"WHERE A0ARTI IN ({','.join('?' * len(grup))})", grup)
        aciklama.update({r[0]: r[1] or '' for r in cu.fetchall()})
    return {'yil': yil, 'ay': ay, 'ts': datetime.now().strftime('%Y-%m-%d %H:%M'),
            'hareket': hareket, 'aciklama': aciklama}


def kaydet(conn, veri, kullanici, sure_sn=None):
    tablolari_kur(conn)
    cur = conn.execute("INSERT INTO aylik_uretim_cekim (yil, ay, ts, kullanici, satir, sure_sn, veri) "
                       "VALUES (?,?,?,?,?,?,?)",
                       (veri['yil'], veri['ay'], veri['ts'], kullanici, len(veri['hareket']), sure_sn,
                        json.dumps(veri, ensure_ascii=False)))
    # Ayın yalnız son 3 okuması tutulur
    conn.execute("DELETE FROM aylik_uretim_cekim WHERE yil=? AND ay=? AND id NOT IN (SELECT id FROM "
                 "aylik_uretim_cekim WHERE yil=? AND ay=? ORDER BY id DESC LIMIT 3)",
                 (veri['yil'], veri['ay'], veri['yil'], veri['ay']))
    conn.commit()
    return cur.lastrowid


def son_cekim(conn, yil, ay):
    tablolari_kur(conn)
    r = conn.execute("SELECT id, ts, kullanici, sure_sn, veri FROM aylik_uretim_cekim WHERE yil=? AND ay=? "
                     "ORDER BY id DESC LIMIT 1", (yil, ay)).fetchone()
    if not r:
        return None
    v = json.loads(r[4])
    v['_id'], v['_kullanici'], v['_sure_sn'] = r[0], r[2], r[3]
    return v


def cekilen_aylar(conn):
    tablolari_kur(conn)
    return [{'yil': r[0], 'ay': r[1], 'ts': r[2]} for r in conn.execute(
        "SELECT yil, ay, MAX(ts) FROM aylik_uretim_cekim GROUP BY yil, ay ORDER BY yil DESC, ay DESC")]


_DENEME = {'calisiyor': False, 'basladi': None, 'bitti': None, 'kim': '', 'hata': '', 'sonuc': None,
           'yil': None, 'ay': None}
_KILIT = threading.Lock()


def deneme_durumu():
    return dict(_DENEME)


def hazirla(conn_ac, erp_ac, yil, ay, kullanici):
    """Senkron: AS400'den oku + sakla."""
    if not _KILIT.acquire(blocking=False):
        raise RuntimeError('Aylık rapor zaten hazırlanıyor')
    t0 = time.time()
    _DENEME.update(calisiyor=True, basladi=datetime.now().strftime('%Y-%m-%d %H:%M:%S'), bitti=None,
                   kim=kullanici, hata='', sonuc=None, yil=yil, ay=ay)
    conn = conn_ac()
    try:
        cn = erp_ac()
        try:
            veri = as400_oku(cn, yil, ay)
        finally:
            try:
                cn.close()
            except Exception:
                pass
        cid = kaydet(conn, veri, kullanici, round(time.time() - t0, 1))
        _DENEME['sonuc'] = {'id': cid, 'satir': len(veri['hareket']),
                            'kod': len({h['kod'] for h in veri['hareket']})}
        return _DENEME['sonuc']
    except Exception as e:
        _DENEME['hata'] = str(e)
        raise
    finally:
        _DENEME.update(calisiyor=False, bitti=datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        try:
            conn.close()
        except Exception:
            pass
        _KILIT.release()


def hazirla_arka_planda(*args):
    if _DENEME['calisiyor'] or _KILIT.locked():
        return False
    _DENEME.update(calisiyor=True, basladi=datetime.now().strftime('%Y-%m-%d %H:%M:%S'), bitti=None, hata='',
                   sonuc=None)

    def _kos():
        try:
            hazirla(*args)
        except Exception as e:
            print(f'[AYLIK ÜRETİM] hazırlama hatası: {e}')
    threading.Thread(target=_kos, daemon=True, name='aylik-uretim').start()
    return True


# ── ÇALIŞMA SAATİ (kapasite Excel'i → 'Çalışma Saati' sayfası) ─────────────
def _bolum_adi(v):
    return ' '.join(str(v or '').strip().upper().split())


def calisma_oku(yol):
    """→ (ay, {excel_bolum: {'nos', 'mesai', 'kisi', 'isimsiz'}}) — İsim · Ay · NçS · %90 · %95 · %100 · BÖLÜM.
    Ayı yazılmamış (isimsiz ek) satırlar dosyadaki ASIL aya sayılır (kapasite modülüyle aynı kural)."""
    from openpyxl import load_workbook
    wb = load_workbook(yol, read_only=True, data_only=True)
    try:
        ad = next((s for s in wb.sheetnames if 'saat' in s.strip().lower()), None)
        if not ad:
            raise ValueError("Dosyada 'Çalışma Saati' sayfası yok — kapasite Excel'ini (Kapasite Kullanım Oranı) yükleyin")
        aylar, out = {}, {}
        for i, r in enumerate(wb[ad].iter_rows(values_only=True), start=1):
            if i == 1 or not r or len(r) < 7:
                continue
            bolum = _bolum_adi(r[6])
            if not bolum:
                continue
            f = lambda v: float(v) if isinstance(v, (int, float)) else 0.0   # noqa: E731
            nos, mesai = f(r[2]), f(r[3]) + f(r[4]) + f(r[5])
            if nos + mesai <= 0:
                continue
            try:
                a = int(float(r[1] or 0))
            except (TypeError, ValueError):
                a = 0
            if a:
                aylar[a] = aylar.get(a, 0) + 1
            d = out.setdefault(bolum, {'nos': 0.0, 'mesai': 0.0, 'kisi': 0, 'isimsiz': 0.0})
            d['nos'] += nos
            d['mesai'] += mesai
            if str(r[0] or '').strip():
                d['kisi'] += 1
            else:
                d['isimsiz'] += nos + mesai
        if not aylar:
            raise ValueError("'Çalışma Saati' sayfasında ay bilgisi bulunamadı")
        return max(aylar, key=aylar.get), out
    finally:
        wb.close()


def database_oku(yol):
    """Kapasite Excel'i 'Database' sayfası → [{kod, saatlik, hat, makine}] (yoksa [])."""
    from openpyxl import load_workbook
    wb = load_workbook(yol, read_only=True, data_only=True)
    try:
        ad = next((s for s in wb.sheetnames if s.strip().lower().startswith('database')), None)
        if not ad:
            return []
        out, gorulen = [], set()
        for i, r in enumerate(wb[ad].iter_rows(values_only=True), start=1):
            if i == 1 or not r or not r[0]:
                continue
            kod = _norm(r[0])
            if kod in gorulen:
                continue
            gorulen.add(kod)
            try:
                saatlik = float(r[2] or 0)
            except (TypeError, ValueError):
                saatlik = 0.0
            out.append({'kod': kod, 'saatlik': saatlik if saatlik > 0 else None,
                        'hat': _bolum_adi(r[3] if len(r) > 3 else ''), 'makine': str(r[1] or '').strip()})
        return out
    finally:
        wb.close()


def database_yaz(conn, satirlar, dosya):
    tablolari_kur(conn)
    if not satirlar:
        return 0
    simdi = datetime.now().strftime('%Y-%m-%d %H:%M')
    conn.execute("DELETE FROM aylik_kapasite_db")
    conn.executemany("INSERT OR REPLACE INTO aylik_kapasite_db (kod, saatlik, hat, makine, dosya, guncellendi) "
                     "VALUES (?,?,?,?,?,?)", [(x['kod'], x['saatlik'], x['hat'], x['makine'],
                                               os.path.basename(str(dosya)), simdi) for x in satirlar])
    conn.commit()
    return len(satirlar)


def calisma_yaz(conn, yil, ay, saatler, dosya, kullanici):
    tablolari_kur(conn)
    conn.execute("DELETE FROM aylik_calisma WHERE yil=? AND ay=?", (yil, ay))
    simdi = datetime.now().strftime('%Y-%m-%d %H:%M')
    conn.executemany("INSERT INTO aylik_calisma (yil, ay, excel_bolum, nos, mesai, kisi, isimsiz, dosya, kullanici, "
                     "guncellendi) VALUES (?,?,?,?,?,?,?,?,?,?)",
                     [(yil, ay, b, round(d['nos'], 1), round(d['mesai'], 1), d['kisi'], round(d['isimsiz'], 1),
                       os.path.basename(str(dosya)), kullanici, simdi) for b, d in saatler.items()])
    conn.commit()
    return len(saatler)


def en_yeni_kapasite(klasor=VARSAYILAN_KAPASITE_KLASORU):
    if not os.path.isdir(klasor):
        raise FileNotFoundError(f'Klasöre erişilemiyor: {klasor}')
    ad = [os.path.join(klasor, f) for f in os.listdir(klasor)
          if f.lower().startswith('kapasite') and f.lower().endswith(('.xlsx', '.xlsm')) and not f.startswith('~$')]
    if not ad:
        raise FileNotFoundError(f"'{klasor}' içinde kapasite Excel'i yok")
    return max(ad, key=os.path.getmtime)


def calisma_aylari(conn):
    tablolari_kur(conn)
    return [{'yil': r[0], 'ay': r[1], 'dosya': r[2]} for r in conn.execute(
        "SELECT yil, ay, MAX(dosya) FROM aylik_calisma GROUP BY yil, ay ORDER BY yil DESC, ay DESC")]


def calisma_haritasi(conn, yil, ay):
    """{(tesis, bölüm): {'nos','mesai','kisi'}} + atanmamış (üretim dışı) satırlar."""
    tablolari_kur(conn)
    atanan, disarida, bilgi = {}, [], None
    for r in conn.execute("SELECT excel_bolum, nos, mesai, kisi, dosya, guncellendi FROM aylik_calisma "
                          "WHERE yil=? AND ay=? ORDER BY excel_bolum", (yil, ay)):
        bilgi = {'dosya': r[4], 'guncellendi': r[5]}
        hedef = CALISMA_BOLUM.get(r[0])
        if hedef:
            d = atanan.setdefault(hedef, {'nos': 0.0, 'mesai': 0.0, 'kisi': 0, 'excel': []})
            d['nos'] += r[1] or 0
            d['mesai'] += r[2] or 0
            d['kisi'] += r[3] or 0
            d['excel'].append(r[0])
        else:
            disarida.append({'excel_bolum': r[0], 'nos': r[1] or 0, 'mesai': r[2] or 0, 'kisi': r[3] or 0})
    return atanan, disarida, bilgi


# ── SINIFLAMA + SÜRE ─────────────────────────────────────────────────────────
def _forge(conn):
    """{tesis: (tam, kök)} — {kod: {bölüm: toplam çevrim sn}}."""
    out = {}
    for lok in ('TK1', 'TK2'):
        tam, kok = {}, {}
        for kod, bolum, ct in conn.execute("SELECT referans_kodu, COALESCE(bolum,'kaynak'), COALESCE(hedef_cycle_time_sn,0) "
                                           "FROM referans_listesi WHERE COALESCE(lokasyon,'TK2')=?", (lok,)):
            for h, k in ((tam, _norm(kod)), (kok, _kok(kod))):
                h.setdefault(k, {}).setdefault(bolum, 0.0)
                h[k][bolum] += float(ct or 0)
        out[lok] = (tam, kok)
    return out


def _sure_kaynaklari(conn):
    kap, rota, db = {}, {}, {}
    try:
        for kod, saatlik, hat in conn.execute("SELECT kod, saatlik, hat FROM aylik_kapasite_db"):
            db[kod] = {'saatlik': saatlik, 'hat': hat or ''}
    except Exception:
        pass
    try:
        for kod, lok, b, sn in conn.execute("SELECT kod, COALESCE(lokasyon,'TK2'), bolum, sure_sn FROM kapasite_sure "
                                            "WHERE sure_sn > 0"):
            kap[(_norm(kod), lok, b)] = float(sn)
    except Exception:
        pass
    try:
        for kod, b, sn in conn.execute("SELECT kod, bolum, sure_sn FROM kapasite_urun_bolum WHERE sure_sn > 0"):
            rota[(_norm(kod), b)] = float(sn)
    except Exception:
        pass
    return kap, rota, db


def rapor(conn, yil, ay):
    """Ayın raporu (son AS400 okumasından). None → o ay hiç okunmamış."""
    import anaveri as AV
    veri = son_cekim(conn, yil, ay)
    if not veri:
        return None
    forge = _forge(conn)
    kap, rota, kdb = _sure_kaynaklari(conn)
    ana = AV.haritasi(conn)
    acik = veri.get('aciklama') or {}
    # Kod × tesis birleştir (causal ayrı tutulur)
    kodlar = {}
    depo_ozet = {}
    for h in veri['hareket']:
        kod = h['kod'].strip()
        n = _norm(kod)
        tesis = DEPO_TESIS.get(h['depo'].upper())
        if not tesis:
            iki = [lok for lok in ('TK2', 'TK1') if forge[lok][0].get(n) or forge[lok][1].get(_kok(kod))]
            tesis = iki[0] if len(iki) == 1 else ''
        d = kodlar.setdefault((n, tesis), {'kod': kod, 'tesis': tesis, 'RPR': 0.0, 'CFI': 0.0, 'depolar': {}})
        d[h['causal']] = d.get(h['causal'], 0.0) + h['adet']
        d['depolar'][h['depo']] = d['depolar'].get(h['depo'], 0.0) + h['adet']
        o = depo_ozet.setdefault(h['depo'] or '?', {'depo': h['depo'] or '?', 'tesis': DEPO_TESIS.get(h['depo'].upper(), ''),
                                                    'RPR': 0.0, 'CFI': 0.0, 'kod': set()})
        o[h['causal']] = o.get(h['causal'], 0.0) + h['adet']
        o['kod'].add(n)
    bolumler, tanimsiz, hammadde, tesis_top = {}, [], [], {}
    for (n, tesis), d in kodlar.items():
        adet = d['RPR'] + d['CFI']
        satir = {'kod': d['kod'], 'aciklama': acik.get(d['kod'], ''), 'tesis': tesis, 'rpr': round(d['RPR'], 2),
                 'cfi': round(d['CFI'], 2), 'adet': round(adet, 2),
                 'depolar': ' '.join(f'{k or "?"}:{v:g}' for k, v in d['depolar'].items())}
        if n.startswith(HAMMADDE_ONEK):
            hammadde.append(satir)
            continue
        # Bölüm: Pull/Pandora fason → Forge (bu tesiste) → kapasite Database → Anaveri → tanımsız
        hedef, kaynak = {}, ''
        a_kod = ana.get(n) or (ana.get(n[:-1]) if n.endswith('W') else None)
        if _fason_mi(a_kod, True):
            tesis = tesis or 'TK1'
            satir['tesis'] = tesis
            hedef, kaynak = {FASON: 0.0}, 'anaveri'
        if tesis and not hedef:
            tam, kok = forge[tesis]
            t_h, k_h = tam.get(n), kok.get(_kok(d['kod'])) or {}
            # Tel: ERP kodu köktür, Forge'da adım ekli satırlar (KESIM, KAPAMA…) — süre kökte
            # TOPLANIR; eksiz satır süresizse kökün toplamı kullanılır.
            hedef = ({b: (sn if sn > 0 else k_h.get(b, 0.0)) for b, sn in t_h.items()} if t_h else dict(k_h))
            kaynak = 'forge' if hedef else ''
        if not hedef:
            # Kapasite Excel'i Database hattı (planlamanın sınıflaması) — tesis uyuşuyorsa
            m = CALISMA_BOLUM.get((kdb.get(n) or {}).get('hat', ''))
            if m and (not tesis or m[0] == tesis):
                tesis = tesis or m[0]
                satir['tesis'] = tesis
                hedef, kaynak = {m[1]: 0.0}, 'kapasite Excel'
        if not hedef and tesis and _fason_mi(a_kod, False):
            hedef, kaynak = {FASON: 0.0}, 'anaveri'
        if not hedef:
            a = a_kod
            m = ANAVERI_BOLUM.get(((a or {}).get('tk', '').upper(), (a or {}).get('hat', '').upper())) if a else None
            if m and (not tesis or m[0] == tesis):
                tesis = tesis or m[0]
                satir['tesis'] = tesis
                hedef, kaynak = {m[1]: 0.0}, 'anaveri'
        if not hedef or not tesis:
            a = ana.get(n)
            satir['anaveri'] = f"{a['tk']} / {a['hat']}" if a else ''
            tanimsiz.append(satir)
            continue
        t = tesis_top.setdefault(tesis, {'adet': 0.0, 'kod': 0, 'rpr': 0.0, 'cfi': 0.0})
        if FASON in hedef:          # tedarikçinin işi tesis üretimine sayılmaz
            t = {'adet': 0.0, 'kod': 0, 'rpr': 0.0, 'cfi': 0.0}
        t['adet'] += adet
        t['kod'] += 1
        t['rpr'] += d['RPR']
        t['cfi'] += d['CFI']
        for b, forge_sn in hedef.items():
            k = kap.get((n, tesis, b)) or kap.get((_kok(d['kod']), tesis, b))
            ks = (kdb.get(n) or {}).get('saatlik')
            if b in SURE_EXCEL_DISI and forge_sn > 0:
                k, ks = None, None
            if b == FASON:
                k, ks, forge_sn = None, None, 0.0
            if k:
                sn, sk = k, 'kapasite Excel'
            elif ks:
                sn, sk = 3600.0 / ks, 'kapasite Excel'
            elif forge_sn > 0:
                sn, sk = forge_sn, 'Forge'
            elif rota.get((n, b)):
                sn, sk = rota[(n, b)], 'ERP rota'
            else:
                sn, sk = 0.0, ''
            bl = bolumler.setdefault((tesis, b), {'tesis': tesis, 'bolum': b, 'ad': BOLUM_AD.get(b, b), 'adet': 0.0,
                                                  'kod': 0, 'rpr': 0.0, 'cfi': 0.0, 'teorik_sn': 0.0,
                                                  'sureli_adet': 0.0, 'satirlar': []})
            bl['adet'] += adet
            bl['kod'] += 1
            bl['rpr'] += d['RPR']
            bl['cfi'] += d['CFI']
            if sn > 0 and adet > 0:
                bl['teorik_sn'] += adet * sn
                bl['sureli_adet'] += adet
            bl['satirlar'].append(dict(satir, sure_sn=round(sn, 2) if sn else None, sure_kaynak=sk,
                                       teorik_saat=round(adet * sn / 3600, 2) if sn and adet > 0 else None,
                                       sinif=kaynak, cok_bolum=len(hedef) > 1))
    calisma, disarida, calisma_bilgi = calisma_haritasi(conn, yil, ay)
    ozet = []
    for tesis in ('TK2', 'TK1'):
        sira = BOLUM_SIRA.get(tesis, ())
        anahtarlar = sorted({b for (t, b) in bolumler if t == tesis} | {b for (t, b) in calisma if t == tesis},
                            key=lambda b: (sira.index(b) if b in sira else 99, b))
        for b in anahtarlar:
            bl = bolumler.get((tesis, b)) or {'tesis': tesis, 'bolum': b, 'ad': BOLUM_AD.get(b, b), 'adet': 0.0,
                                              'kod': 0, 'rpr': 0.0, 'cfi': 0.0, 'teorik_sn': 0.0,
                                              'sureli_adet': 0.0, 'satirlar': []}
            c = calisma.get((tesis, b)) or {}
            nos, mesai = c.get('nos', 0.0), c.get('mesai', 0.0)
            top = nos + mesai
            teorik = bl['teorik_sn'] / 3600
            fsn = b == FASON
            ozet.append({'fason': fsn, 'tesis': tesis, 'bolum': b, 'ad': bl['ad'], 'adet': round(bl['adet'], 2), 'kod': bl['kod'],
                         'rpr': round(bl['rpr'], 2), 'cfi': round(bl['cfi'], 2),
                         'teorik_saat': None if fsn else round(teorik, 1),
                         'sure_kapsami': round(bl['sureli_adet'] / bl['adet'], 3) if bl['adet'] and not fsn else None,
                         'nos': round(nos, 1), 'mesai': round(mesai, 1), 'calisma': round(top, 1),
                         'kisi': c.get('kisi', 0), 'calisma_excel': c.get('excel', []),
                         'performans': round(teorik / top, 3) if top else None,
                         'performans_nos': round(teorik / nos, 3) if nos else None,
                         'satirlar': sorted(bl['satirlar'], key=lambda s: -s['adet'])})
    for t in tesis_top.values():
        for k in ('adet', 'rpr', 'cfi'):
            t[k] = round(t[k], 2)
    tanimsiz.sort(key=lambda s: -s['adet'])
    hammadde.sort(key=lambda s: -s['adet'])
    return {'yil': yil, 'ay': ay, 'ay_ad': f'{AY_AD[ay]} {yil}', 'olcum': veri['ts'], 'olcan': veri.get('_kullanici', ''),
            'devam_eden': (yil, ay) == (date.today().year, date.today().month),
            'ozet': ozet, 'tesis': tesis_top, 'tanimsiz': tanimsiz, 'hammadde': hammadde,
            'depolar': sorted(({**o, 'kod': len(o['kod']), 'RPR': round(o['RPR'], 2), 'CFI': round(o['CFI'], 2)}
                               for o in depo_ozet.values()), key=lambda o: -(o['RPR'] + o['CFI'])),
            'calisma_disarida': disarida, 'calisma_bilgi': calisma_bilgi,
            'tanimsiz_adet': round(sum(s['adet'] for s in tanimsiz), 2),
            'hammadde_adet': round(sum(s['adet'] for s in hammadde), 2)}


# ── EXCEL (yöneticiyle paylaşılacak) ─────────────────────────────────────────
def excel(r):
    import io
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Özet'
    koyu = PatternFill('solid', fgColor='1F3864')
    acik = PatternFill('solid', fgColor='E9E4F7')
    ince = Side(style='thin', color='D0D0D0')
    kenar = Border(top=ince, bottom=ince, left=ince, right=ince)
    ws.append([f"AYLIK ÜRETİM RAPORU — {r['ay_ad']}" + (' (ay devam ediyor)' if r['devam_eden'] else '')])
    ws['A1'].font = Font(bold=True, size=14, color='1F3864')
    ws.append([f"Teyidi verilmiş üretim: AS400 RPR (launch teyidi) + CFI (launch'sız üretim girişi), hareket tarihine "
               f"göre · okuma {r['olcum']} · bölümler Forge tanımlarına göre"])
    ws['A2'].font = Font(italic=True, color='666666')
    ws.append([])
    bas = ['Tesis', 'Bölüm', 'Üretilen adet', 'Kod sayısı', 'RPR adet', 'CFI adet', 'Teorik süre (saat)',
           'Süre kapsamı', 'Çalışma saati (NçS)', 'Mesai (saat)', 'Toplam çalışma (saat)', 'Kişi',
           'Performans', 'Hedef', 'Performans (yalnız NçS)']
    ws.append(bas)
    hb = ws.max_row
    for i in range(1, len(bas) + 1):
        c = ws.cell(row=hb, column=i)
        c.font, c.fill, c.border = Font(bold=True, color='FFFFFF'), koyu, kenar
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ws.row_dimensions[hb].height = 32
    for tesis in ('TK2', 'TK1'):
        satirlar = [o for o in r['ozet'] if o['tesis'] == tesis]
        if not satirlar:
            continue
        ilk = ws.max_row + 1
        for o in satirlar:
            if o.get('fason'):
                continue            # fason satırı tesis toplamının altında ayrı yazılır
            ws.append([tesis, o['ad'], o['adet'], o['kod'], o['rpr'], o['cfi'], o['teorik_saat'], o['sure_kapsami'],
                       o['nos'] or None, o['mesai'] or None, o['calisma'] or None, o['kisi'] or None,
                       o['performans'], HEDEF_PERFORMANS if o['calisma'] else None, o['performans_nos']])
        son = ws.max_row
        t = r['tesis'].get(tesis, {})
        L = get_column_letter
        ws.append([f'{tesis} TOPLAM', '(kod tekil)', t.get('adet'), t.get('kod'), t.get('rpr'), t.get('cfi'),
                   f'=SUM(G{ilk}:G{son})', None, f'=SUM(I{ilk}:I{son})', f'=SUM(J{ilk}:J{son})',
                   f'=SUM(K{ilk}:K{son})', f'=SUM(L{ilk}:L{son})', f'=IF(K{son + 1}>0,G{son + 1}/K{son + 1},"")',
                   HEDEF_PERFORMANS, f'=IF(I{son + 1}>0,G{son + 1}/I{son + 1},"")'])
        for i in range(1, len(bas) + 1):
            c = ws.cell(row=ws.max_row, column=i)
            c.font, c.fill = Font(bold=True), acik
        for o in [x for x in satirlar if x.get('fason')]:
            ws.append([tesis, o['ad'], o['adet'], o['kod'], o['rpr'], o['cfi'], None, None, None, None, None, None,
                       None, None, 'tedarikçinin işi — tesis toplamına ve performansa girmez'])
            ws.cell(row=ws.max_row, column=2).font = Font(italic=True, color='666666')
        for rr in range(ilk, ws.max_row + 1):
            for i in range(1, len(bas) + 1):
                ws.cell(row=rr, column=i).border = kenar
            for col in ('C', 'D', 'E', 'F'):
                ws[f'{col}{rr}'].number_format = '#,##0'
            for col in ('G', 'I', 'J', 'K'):
                ws[f'{col}{rr}'].number_format = '#,##0.0'
            for col in ('H', 'M', 'N', 'O'):
                ws[f'{col}{rr}'].number_format = '0%'
        ws.append([])
    alan = f'M5:M{ws.max_row}'
    ws.conditional_formatting.add(alan, CellIsRule(operator='lessThan', formula=[str(HEDEF_PERFORMANS)],
                                                   fill=PatternFill('solid', fgColor='F8CBAD')))
    ws.conditional_formatting.add(alan, CellIsRule(operator='greaterThanOrEqual', formula=[str(HEDEF_PERFORMANS)],
                                                   fill=PatternFill('solid', fgColor='C6EFCE')))
    notlar = [
        'NOTLAR',
        "• Bir kod tesisinde birden çok bölümde tanımlıysa (ör. kaynak + montaj) her bölümün işine sayılır; tesis toplamı kodu bir kez sayar.",
        '• Teorik süre = adet × birim süre. Birim süre önceliği: kapasite Excel\'i (Database saatlik adet) → Forge çevrim süresi (tel: adımların toplamı) → ERP rota süresi. '
        '"Süre kapsamı" = birim süresi bilinen adetin payı; süresiz kodlar teorik süreye girmez.',
        '• Performans = teorik süre ÷ toplam çalışma saati (NçS + mesai). Son sütun, eski planlama tablosu gibi yalnız NçS\'ye böler.',
        f"• Tanımsız (Forge'da ve Anaveri'de bölümü olmayan) kodlar: {len(r['tanimsiz'])} kod, {r['tanimsiz_adet']:,.0f} adet — "
        "'Tanımsız kodlar' sayfası; özete girmez.",
        f"• Hammadde (20.* / 21.*, metre-kg): {len(r['hammadde'])} kod — adet toplamına girmez.",
        "• Fason (tedarikçi): Anaveri'de Pandora / Pull telleri (Pull şu an fasonda üretiliyor) ve 'Fason' işaretli kodlar "
        "(Forge'da tanımı yoksa) — tesis toplamına ve performansa girmez.",
    ]
    if r.get('calisma_bilgi'):
        notlar.append(f"• Çalışma saati kaynağı: {r['calisma_bilgi']['dosya']} ({r['calisma_bilgi']['guncellendi']}).")
    else:
        notlar.append('• Bu ay için çalışma saati yüklenmedi — performans hesaplanmadı.')
    for n in notlar:
        ws.append([n])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=(n == 'NOTLAR'), color='444444')
    for i, w in enumerate([12, 22, 14, 10, 12, 12, 14, 11, 14, 12, 14, 7, 12, 8, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C5'
    # Bölüm sayfaları
    for o in r['ozet']:
        if not o['satirlar']:
            continue
        ad = f"{o['tesis']} {o['ad']}"[:31].replace('/', '-')
        w = wb.create_sheet(ad)
        w.append([f"{o['tesis']} {o['ad']} — {r['ay_ad']} · {o['kod']} kod · {o['adet']:,.0f} adet · " +
                  ('tedarikçinin işi — tesis toplamına ve performansa girmez' if o.get('fason')
                   else f"teorik {o['teorik_saat']:,.1f} saat")])
        w['A1'].font = Font(bold=True, size=12)
        b2 = ['Kod', 'Açıklama', 'RPR', 'CFI', 'Toplam adet', 'Birim süre (sn)', 'Süre kaynağı', 'Teorik süre (saat)',
              'Sınıflama', 'Depo', 'Not']
        w.append(b2)
        for i in range(1, len(b2) + 1):
            c = w.cell(row=2, column=i)
            c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
        for s in o['satirlar']:
            w.append([s['kod'], s['aciklama'], s['rpr'] or None, s['cfi'] or None, s['adet'], s['sure_sn'],
                      s['sure_kaynak'], s['teorik_saat'], {'forge': 'Forge', 'anaveri': 'Anaveri'}.get(s['sinif'], s['sinif']),
                      s['depolar'], 'birden çok bölümde tanımlı' if s['cok_bolum'] else ''])
        for i, wd in enumerate([18, 34, 10, 10, 12, 12, 14, 13, 10, 18, 22], start=1):
            w.column_dimensions[get_column_letter(i)].width = wd
        w.freeze_panes = 'B3'
        w.auto_filter.ref = f'A2:K{w.max_row}'
    # Tanımsız + hammadde + depo + çalışma saati
    for baslik, liste, kol in (('Tanımsız kodlar', r['tanimsiz'], 'anaveri'), ('Hammadde (adet dışı)', r['hammadde'], None)):
        w = wb.create_sheet(baslik)
        w.append(['Kod', 'Açıklama', 'Tesis (depo)', 'RPR', 'CFI', 'Toplam', 'Depo'] + (['Anaveri (TK-1/2 / Hat)'] if kol else []))
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
        for s in liste:
            w.append([s['kod'], s['aciklama'], s['tesis'], s['rpr'] or None, s['cfi'] or None, s['adet'], s['depolar']]
                     + ([s.get('anaveri', '')] if kol else []))
        for i, wd in enumerate([18, 34, 10, 10, 10, 12, 18, 24], start=1):
            w.column_dimensions[get_column_letter(i)].width = wd
    w = wb.create_sheet('Depo ve çalışma saati')
    w.append(['Depo', 'Tesis', 'RPR', 'CFI', 'Kod sayısı'])
    for c in w[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
    for d in r['depolar']:
        w.append([d['depo'], d['tesis'], d['RPR'], d['CFI'], d['kod']])
    w.append([])
    w.append(['Çalışma saati (kapasite Excel\'i)', 'Forge bölümü', 'NçS', 'Mesai', 'Kişi'])
    for c in w[w.max_row]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
    for o in r['ozet']:
        if o['calisma']:
            w.append([', '.join(o['calisma_excel']), f"{o['tesis']} {o['ad']}", o['nos'], o['mesai'], o['kisi']])
    for d in r['calisma_disarida']:
        w.append([d['excel_bolum'], '— üretim dışı / atanmamış', d['nos'], d['mesai'], d['kisi']])
    for i, wd in enumerate([34, 26, 12, 12, 10], start=1):
        w.column_dimensions[get_column_letter(i)].width = wd
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
