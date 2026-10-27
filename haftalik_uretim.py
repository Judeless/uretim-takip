# -*- coding: utf-8 -*-
"""HAFTALIK ÜRETİM (kullanıcı 2026-10-27): "TK2'de haftalık bazlı geçmişe dönük üretilen ürün
miktarlarını görmek istiyorum — metal enjeksiyon, robot kaynak, montaj. Örnek: W30'da 5000
mekanizma üretilmiş. Hafta bazlı karşılaştırma yapabileyim."

ISO hafta × bölüm OK adet. Kaynak Forge üretim kayıtları (uretim_kayitlari.ok_adet), hafta
VARDİYANIN tarihine göre (gece vardiyası başladığı güne yazılır). ERP teyidi değil — o Aylık
Üretim Raporu'nda (aylik_uretim.py).

YENİ / BİTMİŞ ÜRÜN OLMAYAN KAYITLAR SAYILMAZ — ERP teyidindeki kuralın aynısı (as400/launch_esle):
  · rework kaydı (referans / açıklamada 'rework'),
  · tamir / onarım / punta eki (yeniden işleme),
  · TK1 plastik sızdırmazlık test cihazı (aynı ürün yapıştırmada zaten sayıldı),
  · ARA OPERASYON: 'X 1. op' + 'X 2. op' aynı parçadır — yalnız en son op sayılır; hazırlık ve
    preste op eki taşıyan yazım sayılmaz (op_kurali_uygula; aynı parça iki kez sayılmasın).
Sayılmayan adet hafta × bölüm 'haric' alanında, sebep kırılımı 'haric_ozet'te döner.
Bölümler birbirine TOPLANMAZ: montaj = mekanizma, kaynak = kaynatılan parça, metal = dökülen parça.
"""
import os
import sys
from datetime import date, timedelta

try:
    _d = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'as400')
    if _d not in sys.path:
        sys.path.insert(0, _d)
    import launch_esle as _le
except Exception:                         # kural modülü yoksa hiçbir kayıt dışarıda kalmaz
    _le = None


# Kullanıcının istediği üç bölüm önde (panel ve Excel sırası); diğerleri adede göre
ONCELIK = ('metal', 'kaynak', 'montaj')
SEBEP_AD = {'test': 'test cihazı', 'rework': 'rework', 'tamir': 'tamir / punta', 'ara_op': 'ara operasyon'}


def _yeni_urun_degil(lok, bolum, istasyon, referans, aciklama):
    """Kayıt yeni ürün DEĞİL mi? → kısa sebep | ''."""
    if _le is None:
        return ''
    try:
        if (lok, bolum, _le.makine_adi(lok, bolum, istasyon)) in _le.TEYIT_DISI_MAKINE:
            return 'test'
        if _le.rework_mi(aciklama, referans):
            return 'rework'
        if _le.operator_eki(referans, aciklama):
            return 'tamir'
    except Exception:
        return ''
    return ''


def ref_anahtar(kod):
    """Aynı ürünün farklı yazımları tek satır: '10 300 0609' = '10.300.0609'."""
    if _le is None:
        return (kod or '').strip().upper()
    try:
        return _le.birlestirme_anahtari(_le.bosluk_nokta(kod))
    except Exception:
        return (kod or '').strip().upper()


def hafta_anahtar(gun):
    y, w, _ = gun.isocalendar()
    return f'{y}-W{w:02d}'


def hafta_araligi(anahtar):
    """'2026-W30' → (pazartesi, pazar) date. Hatalıysa ValueError."""
    y, w = anahtar.split('-W')
    bas = date.fromisocalendar(int(y), int(w), 1)
    return bas, bas + timedelta(days=6)


def _kayitlar(conn, lokasyon, bas, bit, haric_bolum=()):
    """Aralığın üretim kayıtları → (sayılan, sayılmayan) dict listeleri
    {tarih, bolum, referans, ok, nok, vid[, sebep]}. Sayılanların referansı son op'ta base koddur."""
    satirlar, disari = [], []
    for t, b, ist, ref, acik, ok, nok, vid in conn.execute(
            "SELECT v.tarih, COALESCE(v.bolum, 'kaynak'), COALESCE(u.istasyon, 0), u.referans_kodu, "
            "COALESCE(u.aciklama, ''), COALESCE(u.ok_adet, 0), COALESCE(u.nok_adet, 0), v.id "
            "FROM uretim_kayitlari u JOIN vardiyalar v ON v.id = u.vardiya_id "
            "WHERE COALESCE(v.lokasyon, 'TK2') = ? AND v.tarih BETWEEN ? AND ?",
            (lokasyon, bas.isoformat(), bit.isoformat())).fetchall():
        if b in haric_bolum:
            continue
        r = {'tarih': t, 'bolum': b, 'referans': (ref or '').strip(), 'ok': ok, 'nok': nok, 'vid': vid}
        sebep = _yeni_urun_degil(lokasyon, b, ist, ref, acik)
        if sebep:
            disari.append(dict(r, sebep=sebep))
        else:
            satirlar.append(r)
    if _le is not None and satirlar:
        try:
            # Aralığın TAMAMI tek seferde: 1. op pazartesi, 2. op salı yapılsa da son op bilinir
            satirlar, ara = _le.op_kurali_uygula(satirlar)
            disari += [dict(r, sebep='ara_op') for r in ara]
        except Exception as e:                 # kural çalışmazsa hepsi sayılır (bilgi konsolda)
            print(f'[haftalik_uretim] op kurali uygulanamadi: {e}')
    return satirlar, disari


def ozet(conn, lokasyon='TK2', n=26, haric_bolum=(), bugun=None):
    """Son `n` ISO hafta (bu hafta dahil; n<=0 → verinin başından) × bölüm.
    → {'lokasyon', 'bolumler': [kod, ... adede göre], 'haftalar': [{anahtar, yil, no, bas, bit,
       devam, b: {bolum: {ok, nok, vardiya, gun, haric}}} ... eskiden yeniye, boş haftalar dahil]}"""
    bugun = bugun or date.today()
    bu_pzt = bugun - timedelta(days=bugun.weekday())
    if n and n > 0:
        bas = bu_pzt - timedelta(weeks=n - 1)
    else:
        ilk = conn.execute("SELECT MIN(v.tarih) FROM vardiyalar v JOIN uretim_kayitlari u ON u.vardiya_id = v.id "
                           "WHERE COALESCE(v.lokasyon, 'TK2') = ?", (lokasyon,)).fetchone()[0]
        bas = date.fromisoformat(ilk) if ilk else bu_pzt
        bas -= timedelta(days=bas.weekday())
    bit = bu_pzt + timedelta(days=6)
    hafta = {}
    vardiya, gunler, haric_ozet = {}, {}, {}
    sayilan, disari = _kayitlar(conn, lokasyon, bas, bit, haric_bolum)
    bos = lambda: {'ok': 0.0, 'nok': 0.0, 'vardiya': 0, 'gun': 0, 'haric': 0.0}   # noqa: E731
    for r in disari:
        a = hafta_anahtar(date.fromisoformat(r['tarih']))
        hafta.setdefault(a, {}).setdefault(r['bolum'], bos())['haric'] += r['ok']
        ho = haric_ozet.setdefault(r['bolum'], {})
        ho[r['sebep']] = round(ho.get(r['sebep'], 0) + r['ok'], 2)
    for r in sayilan:
        a = hafta_anahtar(date.fromisoformat(r['tarih']))
        h = hafta.setdefault(a, {}).setdefault(r['bolum'], bos())
        h['ok'] += r['ok']
        h['nok'] += r['nok']
        vardiya.setdefault((a, r['bolum']), set()).add(r['vid'])
        gunler.setdefault((a, r['bolum']), set()).add(r['tarih'])
    toplam = {}
    for a, bl in hafta.items():
        for b, h in bl.items():
            h['vardiya'] = len(vardiya.get((a, b), ()))
            h['gun'] = len(gunler.get((a, b), ()))
            for k in ('ok', 'nok', 'haric'):
                h[k] = round(h[k], 2)
            toplam[b] = toplam.get(b, 0) + h['ok']
    haftalar = []
    g = bas
    while g <= bit:
        a = hafta_anahtar(g)
        y, w, _ = g.isocalendar()
        haftalar.append({'anahtar': a, 'yil': y, 'no': w, 'bas': g.isoformat(),
                         'bit': (g + timedelta(days=6)).isoformat(), 'devam': g == bu_pzt,
                         'b': hafta.get(a, {})})
        g += timedelta(weeks=1)
    # Baştaki boş haftalar (veri başlamadan önce) gösterilmez; aradaki boş haftalar (tatil) kalır
    while haftalar and not haftalar[0]['b']:
        haftalar.pop(0)
    sira = sorted(toplam, key=lambda b: (ONCELIK.index(b) if b in ONCELIK else len(ONCELIK), -toplam[b]))
    return {'lokasyon': lokasyon, 'bolumler': sira, 'haftalar': haftalar,
            'haric_ozet': haric_ozet, 'sebep_ad': SEBEP_AD}


def karsilastir(conn, lokasyon, a, b, haric_bolum=()):
    """İki ISO hafta, bölüm × referans. a/b: '2026-W30'.
    → {'a': {anahtar, bas, bit}, 'b': {...}, 'bolumler': {bolum: {'a': {ok, vardiya}, 'b': {...},
       'referanslar': [{kod, a, b, fark}] |fark| büyükten}}}"""
    out = {'bolumler': {}}
    for etiket, anahtar in (('a', a), ('b', b)):
        bas, bit = hafta_araligi(anahtar)
        out[etiket] = {'anahtar': anahtar, 'bas': bas.isoformat(), 'bit': bit.isoformat()}
        vardiya = {}
        for x in _kayitlar(conn, lokasyon, bas, bit, haric_bolum)[0]:
            bl = x['bolum']
            d = out['bolumler'].setdefault(bl, {'a': {'ok': 0.0, 'vardiya': 0}, 'b': {'ok': 0.0, 'vardiya': 0},
                                                'ref': {}})
            d[etiket]['ok'] += x['ok']
            vardiya.setdefault(bl, set()).add(x['vid'])
            kod = x['referans'] or '—'
            r = d['ref'].setdefault(ref_anahtar(kod), {'kod': kod, 'a': 0.0, 'b': 0.0})
            if kod.count('.') > r['kod'].count('.'):      # noktalı yazım görünsün ('10 300 0609' değil)
                r['kod'] = kod
            r[etiket] += x['ok']
        for bl, s in vardiya.items():
            out['bolumler'][bl][etiket]['vardiya'] = len(s)
    for bl, d in out['bolumler'].items():
        for k in ('a', 'b'):
            d[k]['ok'] = round(d[k]['ok'], 2)
        refs = [{'kod': r['kod'], 'a': round(r['a'], 2), 'b': round(r['b'], 2), 'fark': round(r['a'] - r['b'], 2)}
                for r in d.pop('ref').values()]
        refs.sort(key=lambda r: (-abs(r['fark']), -max(r['a'], r['b']), r['kod']))
        d['referanslar'] = refs
    return out


def excel(sonuc, bolum_ad, karsi=None):
    """Haftalık tablo (+ verilirse iki hafta karşılaştırması) → .xlsx baytları.
    bolum_ad: {kod: görünen ad}."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as H
    tr = lambda s: '.'.join(reversed(s.split('-')))    # noqa: E731  2026-07-20 → 20.07.2026
    wb = Workbook()
    ws = wb.active
    ws.title = 'Haftalık üretim'
    bl = sonuc['bolumler']
    ws.append([f"Haftalık üretim — {sonuc['lokasyon']} · OK adet (Forge üretim kayıtları; tamir / rework / "
               f"test cihazı ve ara operasyon kayıtları sayılmaz). Bölümler toplanmaz: farklı birimler."])
    ws.append(['Hafta', 'Tarih'] + [x for b in bl for x in (bolum_ad.get(b, b), 'Vardiya')])
    for h in reversed(sonuc['haftalar']):
        satir = [f"W{h['no']:02d} {h['yil']}" + (' (devam ediyor)' if h['devam'] else ''),
                 f"{tr(h['bas'])} – {tr(h['bit'])}"]
        for b in bl:
            x = h['b'].get(b) or {}
            satir += [x.get('ok') or None, x.get('vardiya') or None]
        ws.append(satir)
    ws['A1'].font = Font(bold=True, size=12)
    koyu = PatternFill('solid', fgColor='1F3864')
    for c in ws[2]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, gen in enumerate([20, 24] + [14, 9] * len(bl), start=1):
        ws.column_dimensions[H(i)].width = gen
    for row in ws.iter_rows(min_row=3, min_col=3):
        for c in row:
            c.number_format = '#,##0'
    ws.freeze_panes = 'C3'
    if karsi:
        w2 = wb.create_sheet('Karşılaştırma')
        ka, kb = karsi['a'], karsi['b']
        w2.append([f"{ka['anahtar']} ({tr(ka['bas'])}–{tr(ka['bit'])}) ↔ {kb['anahtar']} ({tr(kb['bas'])}–{tr(kb['bit'])})"])
        w2['A1'].font = Font(bold=True, size=12)
        for b in sorted(karsi['bolumler'], key=lambda x: bl.index(x) if x in bl else 99):
            d = karsi['bolumler'][b]
            w2.append([])
            w2.append([bolum_ad.get(b, b), ka['anahtar'], kb['anahtar'], 'Fark'])
            for c in w2[w2.max_row]:
                c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
            w2.append(['Toplam', d['a']['ok'], d['b']['ok'], round(d['a']['ok'] - d['b']['ok'], 2)])
            w2.append(['Vardiya', d['a']['vardiya'], d['b']['vardiya'], d['a']['vardiya'] - d['b']['vardiya']])
            for c in w2[w2.max_row - 1]:
                c.font = Font(bold=True)
            for r in d['referanslar']:
                w2.append([r['kod'], r['a'] or None, r['b'] or None, r['fark']])
        for i, gen in enumerate([24, 12, 12, 12], start=1):
            w2.column_dimensions[H(i)].width = gen
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
