# -*- coding: utf-8 -*-
"""
anaveri.py — planlamanın ANAVERİ sayfası (PLAN 26xxxx.xlsb), Forge'da ortak tablo.

Sütunlar: CD ART. · PRIORIT. · A/P · TK-1/2 · Hat · Makine · Hammadde · Kapasite 1V/Adet
· zor/kolay. TK-1/2 değerleri: TK-1 · TK-2 · Pandora (tel tedarikçisi) · Tabo (hortum) ·
Satınalma · Fason.

Kullananlar:
  · tel_plani  — telin tedarikçisi (Pandora / TK-1 / Satınalma), Pull kontrolü, kapasite
  · satis_plani — çekimin KAPSAMI (yalnız TK2 mi, TK1+TK2 mi) ve tesis süzgeci

Tohum: tohum/plan_anaveri.json (PLAN 260921, tüm satırlar) — tablo BOŞSA yazılır;
sonrası panelden 'Anaveri yükle' / 'Plan klasöründen oku'.
"""
import json
import os
import re
from datetime import datetime

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
TOHUM = os.path.join(PROJECT_DIR, 'tohum', 'plan_anaveri.json')
VARSAYILAN_PLAN_KLASORU = r'Q:\UretimPlanlama\EMRE\Yeni klasör\Plan'
# TK-1/2 değeri → tesis (satış planı kapsamı için). Pandora telleri ve Tabo hortumları
# TK1 ürünlerinde kullanılır; satınalma/fason kalemleri üretim tesisine sayılmaz.
TESIS = {'TK-1': 'TK1', 'TK1': 'TK1', 'PANDORA': 'TK1', 'TABO': 'TK1', 'TK-2': 'TK2', 'TK2': 'TK2'}

_BASLIK = {'kod': ('CD ART.', 'CD ART'), 'ap': ('A/P',), 'tk': ('TK-1/2', 'TK 1/2'), 'hat': ('HAT',),
           'makine': ('MAKINE', 'MAKİNE'), 'kapasite': ('KAPASITE 1V/ADET', 'KAPASİTE 1V/ADET', 'KAPASITE'),
           'zorluk': ('ZOR/KOLAY', 'ZORLUK')}


def tablolari_kur(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS plan_anaveri (
        kod TEXT PRIMARY KEY, ap TEXT DEFAULT '', tk TEXT DEFAULT '', hat TEXT DEFAULT '',
        makine TEXT DEFAULT '', kapasite REAL, zorluk TEXT DEFAULT '')""")
    conn.execute("""CREATE TABLE IF NOT EXISTS plan_anaveri_yukleme (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, dosya TEXT DEFAULT '',
        satir INTEGER DEFAULT 0, kullanici TEXT DEFAULT '')""")
    # Bizim Ana Veri'mizin 93.* satırları (bkz. TEL_KAYNAK) — plan_anaveri ile aynı şekil
    conn.execute("""CREATE TABLE IF NOT EXISTS ana_veri_tel (
        kod TEXT PRIMARY KEY, ap TEXT DEFAULT '', tk TEXT DEFAULT '', hat TEXT DEFAULT '',
        makine TEXT DEFAULT '', kapasite REAL, zorluk TEXT DEFAULT '')""")
    # 'TK2 mekanizma teli' sütunu (2026-10-13): 1 = işaretli, 0 = boş → tel planında
    # gösterilmez; NULL = dosyada sütun yok (süzme yapılmaz)
    if 'mekanizma' not in {r[1] for r in conn.execute("PRAGMA table_info(ana_veri_tel)")}:
        conn.execute("ALTER TABLE ana_veri_tel ADD COLUMN mekanizma INTEGER")
    conn.execute("""CREATE TABLE IF NOT EXISTS ana_veri_tel_yukleme (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, satir INTEGER DEFAULT 0,
        kullanici TEXT DEFAULT '')""")
    if conn.execute("SELECT COUNT(*) FROM plan_anaveri").fetchone()[0] == 0 and os.path.exists(TOHUM):
        try:
            with open(TOHUM, encoding='utf-8') as f:
                tohum = json.load(f)
            yaz(conn, tohum.get('satirlar') or [], tohum.get('dosya') or 'tohum', 'tohum', commit=False)
        except Exception as e:                    # tohum okunamazsa kullananlar yine çalışır
            print(f'[ANAVERİ] tohum okunamadı: {e}')
    conn.commit()


def _metin(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return '' if s in ('0x2a', '#N/A', '#YOK') else s


def satirlari(satirlar):
    """Ham sayfa satırları → [{kod, ap, tk, hat, makine, kapasite, zorluk}] (tekil kod).
    Başlık satırı ilk 10 satırda 'CD ART' ile aranır; sütunlar ADLA bulunur."""
    satirlar = list(satirlar)
    bas_i, ix = None, {}
    for i, r in enumerate(satirlar[:10]):
        adlar = [_metin(v).upper() for v in r]
        if any(a.startswith('CD ART') for a in adlar):
            bas_i = i
            for alan, adaylar in _BASLIK.items():
                for j, a in enumerate(adlar):
                    if alan not in ix and a in adaylar:
                        ix[alan] = j
            break
    if bas_i is None or 'kod' not in ix or 'tk' not in ix:
        raise ValueError("Anaveri sayfasında 'CD ART.' ve 'TK-1/2' başlıkları bulunamadı")
    out, gorulen = [], set()
    for r in satirlar[bas_i + 1:]:
        def al(a):
            return _metin(r[ix[a]]) if a in ix and ix[a] < len(r) else ''
        kod = re.sub(r'\s+', '', al('kod')).upper()
        if not kod or kod in gorulen or not re.match(r'^[0-9A-Z]', kod):
            continue
        gorulen.add(kod)
        try:
            kap = float(al('kapasite')) if al('kapasite') else None
        except ValueError:
            kap = None
        out.append({'kod': kod, 'ap': al('ap'), 'tk': al('tk'), 'hat': al('hat'), 'makine': al('makine'),
                    'kapasite': kap if kap and kap > 0 else None, 'zorluk': al('zorluk')})
    return out


def dosyadan(yol):
    """PLAN 26xxxx.xlsb / .xlsx → satırlar. Sayfa 'Anaveri' (ya da 'Ana Veri')."""
    if yol.lower().endswith('.xlsb'):
        from pyxlsb import open_workbook
        with open_workbook(yol) as wb:
            ad = next((s for s in wb.sheets if s.replace(' ', '').lower() == 'anaveri'), None)
            if not ad:
                raise ValueError(f"Dosyada 'Anaveri' sayfası yok (sayfalar: {', '.join(wb.sheets)})")
            with wb.get_sheet(ad) as ws:
                return satirlari([c.v for c in r] for r in ws.rows())
    from openpyxl import load_workbook
    wb = load_workbook(yol, read_only=True, data_only=True)
    try:
        ws = next((w for w in wb.worksheets if w.title.replace(' ', '').lower() == 'anaveri'), None)
        if ws is None:
            raise ValueError(f"Dosyada 'Anaveri' sayfası yok (sayfalar: {', '.join(wb.sheetnames)})")
        return satirlari(ws.iter_rows(values_only=True))
    finally:
        wb.close()


def en_yeni_plan(klasor=VARSAYILAN_PLAN_KLASORU):
    """Klasördeki en yeni 'PLAN *.xlsb' (Excel'in kilit/geçici dosyaları hariç)."""
    if not os.path.isdir(klasor):
        raise FileNotFoundError(f'Klasöre erişilemiyor: {klasor}')
    adaylar = [os.path.join(klasor, f) for f in os.listdir(klasor)
               if f.upper().startswith('PLAN ') and f.lower().endswith(('.xlsb', '.xlsx')) and not f.startswith('~$')]
    if not adaylar:
        raise FileNotFoundError(f"'{klasor}' içinde PLAN dosyası yok")
    return max(adaylar, key=os.path.getmtime)


def yaz(conn, satirlar, dosya, kullanici, commit=True):
    if not satirlar:
        raise ValueError("Anaveri'de kod satırı bulunamadı")
    conn.execute("DELETE FROM plan_anaveri")
    conn.executemany("INSERT OR REPLACE INTO plan_anaveri (kod, ap, tk, hat, makine, kapasite, zorluk) "
                     "VALUES (?,?,?,?,?,?,?)",
                     [(s['kod'], s.get('ap', ''), s.get('tk', ''), s.get('hat', ''), s.get('makine', ''),
                       s.get('kapasite'), s.get('zorluk', '')) for s in satirlar])
    conn.execute("INSERT INTO plan_anaveri_yukleme (ts, dosya, satir, kullanici) VALUES (?,?,?,?)",
                 (datetime.now().strftime('%Y-%m-%d %H:%M'), os.path.basename(str(dosya)), len(satirlar), kullanici))
    if commit:
        conn.commit()
    return len(satirlar)


def haritasi(conn, onek=None):
    """{kod: {ap, tk, hat, makine, kapasite, zorluk}} — onek verilirse yalnız o ön ekliler."""
    tablolari_kur(conn)
    sql, par = "SELECT kod, ap, tk, hat, makine, kapasite, zorluk FROM plan_anaveri", ()
    if onek:
        sql, par = sql + " WHERE kod LIKE ?", (onek + '%',)
    return {r[0]: {'ap': r[1] or '', 'tk': r[2] or '', 'hat': r[3] or '', 'makine': r[4] or '',
                   'kapasite': r[5], 'zorluk': r[6] or ''} for r in conn.execute(sql, par)}


def durumu(conn):
    tablolari_kur(conn)
    r = conn.execute("SELECT ts, dosya, satir, kullanici FROM plan_anaveri_yukleme ORDER BY id DESC LIMIT 1").fetchone()
    return dict(zip(('ts', 'dosya', 'satir', 'kullanici'), r)) if r else None


# ── ANA VERİ › teller (kullanıcı 2026-10-13) ───────────────────────────────────
# "Tel planı modülü için kullandığımız referansları da kendi Ana Veri Excel'imize
# ekleyelim; tek Excel'den devam edelim" → "tel referanslarını da Ana Veri sayfasına
# taşıyalım". Planlamanın Anaveri'sinin 93.* satırları bizim 'Ana Veri' sayfasında:
# TK-1/2 (üretici) · Makine (Pandora = Pull teli → hat) · Kapasite 1V/Adet · Zor/Kolay (bkz.
# import_excel.ana_veri_oku). 'Ana Veri Yükle' ile ana_veri_tel tablosuna gelir.
# Tablo DOLUYSA tel planı ve aylık raporun tel sınıflaması YALNIZ buradan okur;
# boşsa (ilk yüklemeye kadar) planlamanın Anaveri'si (plan_anaveri) kullanılır.
TEL_KAYNAK = "Ana Veri › 93.* tel satırları"


def tel_yaz(conn, satirlar, kullanici, commit=True):
    """Ana Veri'nin 93.* satırları → ana_veri_tel (ayna: dosyada olmayan satır silinir).
    commit=False: çağıranın işlemi içinde (tablolar önceden kurulmuş olmalı — kurulum commit eder)."""
    if commit:
        tablolari_kur(conn)
    conn.execute("DELETE FROM ana_veri_tel")
    conn.executemany("INSERT OR REPLACE INTO ana_veri_tel (kod, ap, tk, hat, makine, kapasite, zorluk, mekanizma) "
                     "VALUES (?,?,?,?,?,?,?,?)",
                     [(s['kod'], s.get('ap', ''), s.get('tk', ''), s.get('hat', ''), s.get('makine', ''),
                       s.get('kapasite'), s.get('zorluk', ''), s.get('mekanizma')) for s in satirlar])
    conn.execute("INSERT INTO ana_veri_tel_yukleme (ts, satir, kullanici) VALUES (?,?,?)",
                 (datetime.now().strftime('%Y-%m-%d %H:%M'), len(satirlar), kullanici or ''))
    if commit:
        conn.commit()
    return len(satirlar)


def tel_farki(conn, satirlar, ornek_n=10):
    """Sayfa ↔ ana_veri_tel: {dosyada, eklenen, degisen, cikan, ilk, uretici_bos, ornek_*}.
    ilk=True → tablo boş (ilk yükleme; tel planı bundan sonra Ana Veri'den okur)."""
    alanlar = ('tk', 'hat', 'makine', 'kapasite', 'zorluk')
    mevcut = {r[0]: dict(zip(alanlar + ('mekanizma',), r[1:])) for r in conn.execute(
        "SELECT kod, tk, hat, makine, kapasite, zorluk, mekanizma FROM ana_veri_tel")}
    yeni = {s['kod']: s for s in satirlar}
    eklenen = [k for k in yeni if k not in mevcut]
    cikan = [k for k in mevcut if k not in yeni]
    degisen = []
    for k, s in yeni.items():
        m = mevcut.get(k)
        if m is None:
            continue
        for a in alanlar:
            x, y = s.get(a), m.get(a)
            if a == 'kapasite':
                if (x or None) != (y or None) and not (x and y and abs(float(x) - float(y)) < 0.001):
                    degisen.append(k)
                    break
            elif str(x or '').strip() != str(y or '').strip():
                degisen.append(k)
                break
    bos = sorted(k for k, s in yeni.items() if not str(s.get('tk') or '').strip())
    # 'TK2 mekanizma teli' işareti ayrı sayılır (yalnız tel planının süzgeci)
    mek_degisen = sorted(k for k, s in yeni.items() if k in mevcut and s.get('mekanizma') != mevcut[k]['mekanizma'])
    return {'dosyada': len(yeni), 'mevcut': len(mevcut), 'eklenen': len(eklenen), 'degisen': len(degisen),
            'cikan': len(cikan), 'ilk': not mevcut, 'uretici_bos': len(bos),
            'mekanizma_degisen': len(mek_degisen), 'ornek_mekanizma_degisen': mek_degisen[:ornek_n],
            'mekanizma_isaretli': (sum(1 for s in yeni.values() if s.get('mekanizma') == 1)
                                   if any(s.get('mekanizma') is not None for s in yeni.values()) else None),
            'ornek_eklenen': sorted(eklenen)[:ornek_n], 'ornek_degisen': sorted(degisen)[:ornek_n],
            'ornek_cikan': sorted(cikan)[:ornek_n], 'ornek_uretici_bos': bos[:ornek_n]}


def tel_kaynagi_ana_veri(conn):
    tablolari_kur(conn)
    return conn.execute("SELECT 1 FROM ana_veri_tel LIMIT 1").fetchone() is not None


def tel_haritasi(conn):
    """Tel bilgisi (93.*): Ana Veri'nin tel satırları yüklendiyse YALNIZ onlar; değilse
    planlamanın Anaveri'si. Biçim haritasi() ile aynı."""
    if not tel_kaynagi_ana_veri(conn):
        return haritasi(conn, '93.')
    return {r[0]: {'ap': r[1] or '', 'tk': r[2] or '', 'hat': r[3] or '', 'makine': r[4] or '',
                   'kapasite': r[5], 'zorluk': r[6] or '', 'mekanizma': r[7]}
            for r in conn.execute("SELECT kod, ap, tk, hat, makine, kapasite, zorluk, mekanizma FROM ana_veri_tel")}


def tel_durumu(conn):
    """Tel bilgisinin kaynağı + son yükleme → {kaynak: 'ana_veri'|'plan', ts, dosya, satir, kullanici}."""
    if tel_kaynagi_ana_veri(conn):
        r = conn.execute("SELECT ts, satir, kullanici FROM ana_veri_tel_yukleme ORDER BY id DESC LIMIT 1").fetchone()
        n = conn.execute("SELECT COUNT(*) FROM ana_veri_tel").fetchone()[0]
        return {'kaynak': 'ana_veri', 'dosya': TEL_KAYNAK, 'satir': n,
                'ts': r[0] if r else '', 'kullanici': r[2] if r else ''}
    d = durumu(conn)
    return dict(d, kaynak='plan') if d else None
