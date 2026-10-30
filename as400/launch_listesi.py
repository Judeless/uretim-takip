# -*- coding: utf-8 -*-
"""DENETİM — açılan / kapatılan launch (iş emri) listesi (kullanıcı 2026-10-30).

İç kontrol talebi madde 1: "Açılan ve kapatılan tüm launch / iş emirlerinin listesi: iş emri
no, ürün kodu, açılış ve kapanış tarihi, planlanan / üretilen / kapatılan adet, kapatan
kullanıcı, finalize edildiği yer (TK1 / TK2)" — 15–19.06.2026 ve 07–11.09.2026 (ADP istenmedi).

SALT OKUNUR — yalnız SELECT. SUNUCUDA, promanage (RDP) oturumunda çalışır (ODBC sürücüsü ve
şifre kasası orada; laptoptan AS400'e bağlanılmaz):
    sunucu\\Launch_Listesi.bat
    python as400\\launch_listesi.py [2026-06-15:2026-06-19 2026-09-07:2026-09-11]

KOLONLAR — 1. çalıştırmanın (2026-10-30 09:38) SYSCOLUMNS + ham verisiyle doğrulandı:
  BPROF0 (OP ana dosyası, kapalı emirler dahil) — ERP tarihi 4 kolon: SS (20) · AA · AY · GÜN
    Q0LPD1-4 'Dt inizio lancio' = LAUNCH (AÇILIŞ) tarihi — 309 emrin 306'sında ilk teyit ≥ bu tarih
    Q0ORD1-4 'Dt ordine'        = sipariş (OPR) tarihi; launch 0–6 gün sonra
    Q0IPD/Q0FPD                  = PLANLANAN üretim başı / sonu (açılış/kapanış DEĞİL)
    Q0DAT1-4 + Q0SER 'ult. variaz.' = kartın son ELLE değişikliği — KAPANIŞ DEĞİL (çoğu
                                   emirde launch tarihinden bile önce); Q0LNCV hep boş
    Q0MGPR ana depo → tesis: 01D/MT2 = TK2 · 01/02/02B/02Z = TK1 · 02M = Aspanda fason
           (1. çalıştırmada Forge ürün eşlemesiyle çapraz kontrol: 02M'yi 54/57 ASPANDA açmış)
  BPROF0'da KAPANIŞ TARİHİ YOK. Kapanış = ERP teyit başlığı BPRCF0'daki SALDO satırı
  (J0FLSA='S'; bizim BPRCF0I import tablosunun aslı), kullanıcı J0SER. TARİH: saldo satırı bir
  depo hareketine bağlıysa (J0MGSS/AA/NU/PG = MGSERE/MGANRE/MGNURE/MGPRRE) o hareketin KAYIT
  tarihi; 0 adetlik kapatmanın hareketi yok → J0DAT ('dt ult. variaz.' — 24 satırda hareketten
  1–30 gün SONRA: satır sonradan değiştirilmiş).
  2. çalıştırma (10:12) kanıtı: kapalı emirlerin HEPSİNDE saldo satırı var; çoğu launch SONRADAN
  ayrı bir 0 adetlik saldo satırıyla kapatılıyor (ör. 26-77711: teyit 17.06 KORAYTK, kapanış
  02.09 ONURTK1). 0 adetlik satırın depo hareketi YOK → son RPR'yi kapanış saymak YANLIŞTI.
  Saldo satırı yoksa (BPRCF0 okunamazsa) yedek: son RPR hareketi + MGUTVA.
  BPRCF0'da aynı numara başka yıllarda da var (J0RED2 18…26) → eşleşme HER ZAMAN yıl + numara.
"""
import io
import os
import re
import sqlite3
import sys
from datetime import date, datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
AS = os.path.dirname(os.path.abspath(__file__))
KOK = os.path.dirname(AS)
sys.path.insert(0, AS)
sys.path.insert(0, KOK)
import as400_config as CFG  # noqa: E402

SEMA = 'TKC0301F'
DONEMLER = [('2026-06-15', '2026-06-19'), ('2026-09-07', '2026-09-11')]
SAYISAL = ('DECIMAL', 'NUMERIC', 'INTEGER', 'SMALLINT', 'BIGINT')
ACILIS = 'Q0LPD'          # Dt inizio lancio
SIPARIS = 'Q0ORD'         # Dt ordine
DEPO_TESIS = {'01D': 'TK2', 'MT2': 'TK2', '01': 'TK1', '02': 'TK1', '02B': 'TK1', '02Z': 'TK1',
              '02M': 'Fason (Aspanda)'}
DURUM_AD = {10: 'OPR (launch alınmamış)', 40: 'Launch açık', 45: 'Launch açık (TK1 akışı)',
            50: 'Launch açık (TK1 akışı)', 70: 'Kapalı'}
_TR = str.maketrans('"£$§°', 'ÜÖİŞÇ')     # EBCDIC(1026)→ODBC bozulması: AL"M$NYUM PAR°A
MG_TARIH = '(MGDSSO*1000000 + MGDAAO*10000 + MGDMMO*100 + MGDGGO)'


def _t(v):
    return '' if v is None else str(v).strip()


def _k(v):
    """Numara parçası anahtarı: Decimal('26') / 26 / '26' → '26'."""
    t = _t(v)
    return t[:-2] if t.endswith('.0') else t


_KONTROL = re.compile('[' + ''.join(chr(c) for c in list(range(0, 9)) + [11, 12] + list(range(14, 32))) + ']')


def _hucre(v):
    """Excel'e yazılabilir değer: Decimal → float, bytes → hex, kontrol karakteri atılır."""
    if hasattr(v, 'as_tuple'):
        return float(v)
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    if isinstance(v, str):
        return _KONTROL.sub('', v).strip()
    return v


def _sayi(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _gun4(*p):
    """(SS, AA, AY, GÜN) → date | None."""
    try:
        v = [int(_sayi(x) or 0) for x in p]
        return date(v[0] * 100 + v[1], v[2], v[3]) if v[2] and v[3] else None
    except (TypeError, ValueError):
        return None


def kolonlar(cur, tablo):
    cur.execute("SELECT COLUMN_NAME, DATA_TYPE, LENGTH, NUMERIC_SCALE, COALESCE(COLUMN_TEXT, ''), "
                "COALESCE(COLUMN_HEADING, '') FROM QSYS2.SYSCOLUMNS "
                "WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ? ORDER BY ORDINAL_POSITION", (SEMA, tablo))
    return [{'ad': _t(r[0]), 'tip': _t(r[1]).upper(), 'uzunluk': r[2], 'olcek': r[3] or 0,
             'aciklama': _t(r[4]), 'baslik': re.sub(r'\s+', ' ', _t(r[5]))} for r in cur.fetchall()]


def _metin(k):
    return ' / '.join(dict.fromkeys(x.upper() for x in (k['aciklama'], k['baslik']) if x))


def tesis_haritasi():
    """Forge'un ürün → TK1/TK2 eşlemesi — yalnız ana deposu bilinmeyen emir için yedek."""
    try:
        import satis_plani as SP
        c = sqlite3.connect(f"file:{os.path.join(KOK, 'uretim.db')}?mode=ro", uri=True)
        try:
            h = SP.tesis_haritasi(c)
        finally:
            c.close()
        return h, SP.tesis_bul
    except Exception as e:
        print(f'  (Forge ürün eşlemesi okunamadı: {e})')
        return {}, lambda kod, h: ''


def _parcali(cur, sql_sablon, degerler, n=400):
    """IN (...) sorgusunu n'erli parçalarla çalıştırır → satırlar (+ description)."""
    out, aciklama = [], None
    for i in range(0, len(degerler), n):
        parca = degerler[i:i + n]
        cur.execute(sql_sablon.format(yer=','.join('?' * len(parca))), parca)
        aciklama = cur.description
        out += cur.fetchall()
    return out, aciklama


def main(donemler):
    print('Açılan / kapatılan launch listesi — salt okunur')
    print('Dönemler: ' + ', '.join(f'{a} – {b}' for a, b in donemler))
    araliklar = [(int(a.replace('-', '')), int(b.replace('-', ''))) for a, b in donemler]

    def donemde(d):
        if not d:
            return ''
        n = d.year * 10000 + d.month * 100 + d.day
        return next((f'{a[8:10]}.{a[5:7]}–{b[8:10]}.{b[5:7]}.{b[:4]}'
                     for (a, b), (x, y) in zip(donemler, araliklar) if x <= n <= y), '')

    def aralik_sql(ifade):
        return '(' + ' OR '.join(f'{ifade} BETWEEN {a} AND {b}' for a, b in araliklar) + ')'

    cn = CFG.baglan(timeout=30)
    cur = cn.cursor()
    print(f'AS400 bağlantısı OK (profil {CFG.DB_KULLANICI})')
    kol_m = kolonlar(cur, 'BMMAF0')
    m_adlar = {k['ad'] for k in kol_m}
    renu_sayisal = next((k['tip'] in SAYISAL for k in kol_m if k['ad'] == 'MGRENU'), True)
    nu = lambda x: int(x) if renu_sayisal and x.isdigit() else x       # noqa: E731
    anahtar = lambda red2, renu: (_k(red2), _k(renu))                     # noqa: E731

    # ── 1) Adaylar: dönemde LAUNCH edilen + dönemde teyit hareketi (RPR) olan emirler ──
    lp = f'({ACILIS}1*1000000 + {ACILIS}2*10000 + {ACILIS}3*100 + {ACILIS}4)'
    secim_p = f"SELECT p.*, a.A0ARDS AS URUN_TANIMI FROM {SEMA}.BPROF0 p LEFT JOIN {SEMA}.BARTF0 a ON a.A0ARTI = p.Q0ARTI"
    cur.execute(f'{secim_p} WHERE {aralik_sql(lp)}')
    adlar = [d[0].strip().upper() for d in cur.description]
    emirler = {}
    for r in cur.fetchall():
        d = dict(zip(adlar, r))
        emirler[anahtar(d.get('Q0RED2'), d.get('Q0RENU'))] = d
    print(f'Dönemde launch edilen: {len(emirler)}')
    # Dönemde KAPANAN: BPRCF0 saldo satırı dönemde (0 adetle kapatmanın depo hareketi yok);
    # BPRCF0 okunamazsa yedek olarak dönemde RPR hareketi olanlar.
    kol_c, bprc_not = [], ''
    try:
        kol_c = kolonlar(cur, 'BPRCF0')
    except Exception as e:
        bprc_not = f'kolonlar okunamadı: {e}'
    bprc = {'J0RED2', 'J0RENU', 'J0FLSA', 'J0SER', 'J0DAT1', 'J0DAT4'} <= {k['ad'] for k in kol_c}
    cur.execute(f"SELECT DISTINCT MGRED2, MGRENU FROM {SEMA}.BMMAF0 WHERE MGCACD = 'RPR' AND {aralik_sql(MG_TARIH)}")
    ek = {anahtar(r[0], r[1]) for r in cur.fetchall()}
    if bprc:
        jd = '(J0DAT1*1000000 + J0DAT2*10000 + J0DAT3*100 + J0DAT4)'
        cur.execute(f"SELECT DISTINCT J0RED2, J0RENU FROM {SEMA}.BPRCF0 WHERE J0FLSA = 'S' AND {aralik_sql(jd)}")
        ek |= {anahtar(r[0], r[1]) for r in cur.fetchall()}
    else:
        bprc_not = bprc_not or 'BPRCF0 yok ya da J0RED2 / J0RENU / J0FLSA / J0SER / J0DAT kolonu yok'
    ek -= set(emirler)
    satir, _ = _parcali(cur, secim_p + ' WHERE p.Q0RENU IN ({yer})', [nu(x) for x in sorted({e[1] for e in ek})])
    for r in satir:
        d = dict(zip(adlar, r))
        if anahtar(d.get('Q0RED2'), d.get('Q0RENU')) in ek:
            emirler[anahtar(d.get('Q0RED2'), d.get('Q0RENU'))] = d
    print(f'Dönemde teyidi ya da saldo (kapanış) kaydı olan ek emir: {len(ek)} · toplam aday: {len(emirler)}')
    renular = [nu(x) for x in sorted({k[1] for k in emirler})]

    # ── 2) Teyit hareketleri (BMMAF0 RPR) ──
    m_ek = [c for c in ('MGSERE', 'MGANRE', 'MGNURE', 'MGPRRE', 'MGUTVA') if c in m_adlar]
    satir, _ = _parcali(cur, f"SELECT MGRED2, MGRENU, MGDSSO, MGDAAO, MGDMMO, MGDGGO, MGQTA{''.join(', ' + c for c in m_ek)} "
                             f"FROM {SEMA}.BMMAF0 WHERE MGCACD = 'RPR' AND MGRENU IN ({{yer}})", renular)
    hareket = {}            # launch → [{tarih, adet, kul, reg}]
    for r in satir:
        ekv = dict(zip(m_ek, r[7:]))
        reg = tuple(_k(ekv.get(c)) for c in ('MGSERE', 'MGANRE', 'MGNURE', 'MGPRRE'))
        hareket.setdefault(anahtar(r[0], r[1]), []).append({
            'tarih': _gun4(*r[2:6]), 'adet': _sayi(r[6]) or 0, 'kul': _t(ekv.get('MGUTVA')), 'reg': reg})
    print(f'Teyit hareketi: {len(satir)} (launch: {len(hareket)})')
    hareket_reg = {h['reg']: h for l in hareket.values() for h in l}

    # ── 3) Saldo (kapanış) kaydı: ERP teyit başlığı BPRCF0 — tarih J0DAT, kullanıcı J0SER ──
    saldo, ham_c, c_adlar = {}, [], []
    if bprc:
        try:
            c_sayisal = next((k['tip'] in SAYISAL for k in kol_c if k['ad'] == 'J0RENU'), True)
            ham_c, aciklama = _parcali(cur, f'SELECT * FROM {SEMA}.BPRCF0 WHERE J0RENU IN ({{yer}})',
                                       [int(x) if c_sayisal else str(x) for x in renular])
            c_adlar = [d[0].strip().upper() for d in (aciklama or [])]
            i2, inu = c_adlar.index('J0RED2'), c_adlar.index('J0RENU')
            ham_c = [r for r in ham_c if anahtar(r[i2], r[inu]) in emirler]      # başka yılın aynı numarası
            for r in ham_c:
                d = dict(zip(c_adlar, r))
                if _t(d.get('J0FLSA')).upper() != 'S':
                    continue
                key = anahtar(d.get('J0RED2'), d.get('J0RENU'))
                h = hareket_reg.get(tuple(_k(d.get(c)) for c in ('J0MGSS', 'J0MGAA', 'J0MGNU', 'J0MGPG')))
                yeni = {'tarih': (h or {}).get('tarih') or _gun4(*(d.get(f'J0DAT{i}') for i in range(1, 5))),
                        'kul': _t(d.get('J0SER')) or (h or {}).get('kul', ''), 'hareketli': bool(h),
                        'adet': _sayi(d.get('J0QTRI')) or 0, 'sira': _sayi(d.get('J0CNPR')) or 0}
                onceki = saldo.get(key)
                if not onceki or ((yeni['tarih'] or date.min, yeni['sira']) >=
                                  (onceki['tarih'] or date.min, onceki['sira'])):
                    saldo[key] = yeni
            bprc_not = f'{len(ham_c)} teyit satırı, {len(saldo)} emirde saldo (S)'
        except Exception as e:                # okunamazsa son RPR yedeği
            bprc_not = f'okunamadı: {e}'
    print(f'BPRCF0: {bprc_not}')

    # ── 4) Dönemde teyidi olup emri bulunamayanlar için ek okuma yok — liste ──
    cn.close()
    harita, tesis_bul = tesis_haritasi()
    satirlar = []
    for key, r in emirler.items():
        durum = int(_sayi(r.get('Q0AVAN')) or -1)
        hl = sorted(hareket.get(key, []), key=lambda h: (h['tarih'] or date.min, h['reg']))
        son = hl[-1] if hl else {}
        ac = _gun4(*(r.get(f'{ACILIS}{i}') for i in range(1, 5)))
        sip = _gun4(*(r.get(f'{SIPARIS}{i}') for i in range(1, 5)))
        kp, kapatan, kaynak = None, '', ''
        if durum == 70:
            sd = saldo.get(key)
            if sd and sd['tarih']:
                kp, kapatan = sd['tarih'], sd['kul']
                kaynak = ('Saldo kaydı (son teyitle birlikte)' if sd['hareketli'] else
                          'Saldo kaydı (0 adet — ayrı kapatma kaydı)' if not sd['adet'] else
                          'Saldo kaydı (hareketi bulunamadı — kayıt tarihi)')
            elif son.get('tarih'):
                kp, kapatan, kaynak = son['tarih'], son['kul'], 'Saldo kaydı yok — son teyit (BMMAF0)'
            else:
                kaynak = 'Kapanış kaydı bulunamadı'
        d_ac, d_kp = donemde(ac), donemde(kp)
        if not (d_ac or d_kp):
            continue                          # ne launch ne kapanış dönemde
        kod = _t(r.get('Q0ARTI'))
        depo = _t(r.get('Q0MGPR'))
        tesis = DEPO_TESIS.get(depo) or ('TK1' if durum in (45, 50) else tesis_bul(kod, harita)) or '?'
        planlanan = _sayi(r.get('Q0QTOR')) or 0
        uretilen = _sayi(r.get('Q0QTRI')) or 0
        donem_adet = {}
        for h in hl:
            if donemde(h['tarih']):
                donem_adet[donemde(h['tarih'])] = donem_adet.get(donemde(h['tarih']), 0) + h['adet']
        # Haziran'da açılıp Eylül'de kapanan emir İKİ dönemin listesinde de görünür (ayrı satır)
        for dn in dict.fromkeys(x for x in (d_ac, d_kp) if x):
            olay = ' + '.join(x for x, v in (('Açıldı', d_ac), ('Kapandı', d_kp)) if v == dn)
            satirlar.append([_hucre(x) for x in [
                dn, olay, f'{key[0]}-{key[1]}', kod, _t(r.get('URUN_TANIMI')).translate(_TR), tesis, depo,
                sip, ac, kp, f'{durum} · {DURUM_AD.get(durum, "?")}', planlanan, uretilen,
                uretilen if durum == 70 else None, kapatan, kaynak,
                hl[0]['tarih'] if hl else None, son.get('tarih'), len(hl), donem_adet.get(dn, 0)]])
    satirlar.sort(key=lambda s: (s[0], s[8] or s[9] or date.min, s[2]))
    print(f'Listeye giren satır: {len(satirlar)}')

    # ── 5) Excel ──
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as H
    wb = Workbook()
    ws = wb.active
    ws.title = 'Launch listesi'
    koyu = PatternFill('solid', fgColor='1F3864')

    def baslik(sayfa, satir_no=1):
        for c in sayfa[satir_no]:
            c.font, c.fill = Font(bold=True, color='FFFFFF'), koyu
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

    ws.append(['Dönem', 'Dönemdeki olay', 'İş emri (launch) no', 'Ürün kodu', 'Ürün tanımı', 'Tesis', 'Ana depo',
               'Sipariş (OPR) tarihi', 'Açılış (launch) tarihi', 'Kapanış tarihi', 'Durum (ERP)',
               'Planlanan adet', 'Üretilen (teyit) adet', 'Kapatılan adet', 'Kapatan kullanıcı', 'Kapanış kaynağı',
               'İlk teyit', 'Son teyit', 'Teyit sayısı', 'Dönemde teyit edilen'])
    baslik(ws)
    for s in satirlar:
        ws.append(s)
    for i, gen in enumerate([17, 16, 13, 18, 34, 15, 8, 12, 12, 12, 22, 11, 11, 11, 14, 30, 11, 11, 9, 11], start=1):
        ws.column_dimensions[H(i)].width = gen
    for row in ws.iter_rows(min_row=2):
        for i in (7, 8, 9, 16, 17):
            row[i].number_format = 'DD.MM.YYYY'
        for i in (11, 12, 13, 19):
            row[i].number_format = '#,##0'
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{H(ws.max_column)}{max(1, ws.max_row)}'
    ws.row_dimensions[1].height = 32

    oz = wb.create_sheet('Özet')
    oz.append(['Dönem', 'Tesis', 'Açılan launch', 'Kapanan launch', 'Kapanan — planlanan adet', 'Kapanan — üretilen adet'])
    baslik(oz)
    toplam = {}
    for s in satirlar:
        t = toplam.setdefault((s[0], s[5]), [0, 0, 0.0, 0.0])
        if 'Açıldı' in s[1]:
            t[0] += 1
        if 'Kapandı' in s[1]:
            t[1] += 1
            t[2] += s[11] or 0
            t[3] += s[12] or 0
    for (d, tesis), t in sorted(toplam.items()):
        oz.append([d, tesis] + t)
    teyitsiz = sum(1 for s in satirlar if s[15] == 'Kapanış kaydı bulunamadı')
    oz.append([])
    for x in ['Açılış = launch tarihi (BPROF0 Q0LPD "Dt inizio lancio"). Sipariş tarihi = OPR (Q0ORD).',
              f"Kapanış = ERP teyit başlığındaki SALDO kaydı (BPRCF0, J0FLSA='S'), kullanıcı J0SER ({bprc_not}). "
              "Tarih: saldo teyidinin depo hareketi (BMMAF0) kayıt tarihi; 0 adetlik kapatmada kaydın tarihi (J0DAT).",
              "Launch'ların çoğu son teyitten sonra AYRI bir 0 adetlik saldo kaydıyla kapatılıyor — kapanış tarihi o kayıttır.",
              'İlk / son teyit = üretim teyit hareketleri (BMMAF0 RPR).',
              'Tesis = emrin ana deposu (Q0MGPR): 01D / MT2 = TK2 · 01 / 02 / 02B / 02Z = TK1 · 02M = Aspanda fason.',
              f"Kapalı olup kapanış kaydı bulunamayan: {teyitsiz} satır."]:
        oz.append([x])
    for i, gen in enumerate([22, 16, 14, 14, 22, 22], start=1):
        oz.column_dimensions[H(i)].width = gen

    for ad, kol in (('BMMAF0 kolonları', kol_m), ('BPRCF0 kolonları', kol_c)):
        if not kol:
            continue
        w = wb.create_sheet(ad)
        w.append(['Kolon', 'Tip', 'Uzunluk', 'Ondalık', 'Açıklama', 'Başlık'])
        baslik(w)
        for k in kol:
            w.append([_hucre(x) for x in (k['ad'], k['tip'], k['uzunluk'], k['olcek'], k['aciklama'], k['baslik'])])
        for i, gen in enumerate([12, 10, 8, 8, 40, 40], start=1):
            w.column_dimensions[H(i)].width = gen
    for ad, bas, liste in (('Ham BPROF0', adlar, [list(r.values()) for r in emirler.values()]),
                           ('Ham BPRCF0', c_adlar, ham_c)):
        if not bas:
            continue
        hw = wb.create_sheet(ad)
        hw.append(bas)
        baslik(hw)
        for r in liste:
            hw.append([_hucre(v) for v in r])
        hw.freeze_panes = 'A2'

    ad = f"Launch_Listesi_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    klasor = os.path.join(KOK, 'data', 'denetim')
    os.makedirs(klasor, exist_ok=True)
    yol = os.path.join(klasor, ad)
    wb.save(yol)
    print(f'\nKaydedildi: {yol}')
    masa = os.path.join(os.path.expanduser('~'), 'Desktop')
    if os.path.isdir(masa):
        import shutil
        shutil.copy2(yol, os.path.join(masa, ad))
        print(f'Masaüstüne de kopyalandı: {os.path.join(masa, ad)}')
    return yol


if __name__ == '__main__':
    donem = []
    for a in sys.argv[1:]:
        m = re.match(r'^(\d{4}-\d{2}-\d{2}):(\d{4}-\d{2}-\d{2})$', a)
        if not m:
            print(f'Dönem biçimi YYYY-AA-GG:YYYY-AA-GG olmalı: {a}')
            sys.exit(2)
        donem.append((m.group(1), m.group(2)))
    try:
        main(donem or DONEMLER)
    except Exception as e:
        print(f'\nHATA: {e!r}')
        if getattr(CFG, 'kimlik_hatasi_mi', None) and CFG.kimlik_hatasi_mi(e):
            print('  -> KİMLİK HATASI: tekrar denemeyin (profil kilitlenir); İtalya IT ile kontrol edin.')
        sys.exit(1)
