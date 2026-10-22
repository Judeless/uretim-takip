# -*- coding: utf-8 -*-
"""BÖLÜM KAPASİTESİ — Ana Veri'ye göre 4 haftalık yük / kişi analizi (kullanıcı 2026-10-21).

Kullanıcı: "ana verimize göre bölüm bazlı kapasite analizi yapalım. 4 haftalık kapasite
bakalım. bölümde çalışacak operatör sayılarını elle düzeltebileyim anlık. mevcuttaki
modülü beğenmedim ve eski database ile yapmıştık. daha doğru ve kolay anlaşılabilir olan
bir şey yapalım sonrasında onu görselleştirelim de grafiklerle."

HESAP (tek ölçü: KİŞİ)
  İhtiyaç   AS400 açık iş emirleri (XPRO90: OPR + açık launch, kalan = adet − teyit) —
            kapasite.talep_senkron'un kapasite_talep tablosu. Emir bitiş tarihinin haftasına
            yazılır (bitişi PAZARTESİ olan bir önceki haftaya); bu haftanın pazartesisinden
            eskisi 'gecikmiş' (bakiye, ayrı sütun).
  Bölüm     Ana Veri (data/AnaVeri.xlsx): D 'TK-1/2' + E 'Bölüm'. TK2 bölümleri Forge
            anahtarıyla (Kaynak, Lazer Kesim, Büküm, Metal Enjeksiyon, İşleme, Montaj), diğer
            tesisler E'deki adla (TK1 Tel, LF, Joystick…). Satın alma / ADP / Aspanda dışarıda.
  Süre      G 'Çevrim süresi / söktak' (sn/adet). G yoksa M 'Kapasite 1V/Adet' (bir kişinin bir
            vardiyada yaptığı adet). METAL: G baskı başı çevrim → parça başı = G ÷ I 'Kalıp göz'
            (kullanıcı 2026-10-22: 3914W 4 göz 65 sn, 811 adet = 3,66 saat). KAYNAK: çift istasyon —
            robot bir istasyonda kaynak yaparken operatör diğerinde söktak yapar (oee.pair_cycle):
            parça başı = K + S − uyum × min(K, S); uyum %100 → max(K,S)'ye iner, %0 → seri (K+S).
            Uyum bölüm ayarı (varsayılan %50, kullanıcı 2026-10-22); min(K,S) payı ayrı taşınır,
            uyum panelde değişince sayfa sunucuya gitmeden yeniden hesaplanır.
  Birim     Kişi ya da MAKİNE (lazer: 6 kurulu makinenin 4'ü aynı anda çalışıyor → 'aktif makine';
            kurulu sayı 'devreye alınabilecek ek makine' için).
  Kişi-saat sn'li ürün: adet × sn ÷ 3600 ÷ verim · vardiya adetli ürün: adet ÷ M × günlük saat
            (M zaten gerçekçi vardiya adedi → verim ikinci kez düşülmez)
  Gereken kişi = haftalık kişi-saat ÷ (o haftanın ÇALIŞMA GÜNÜ × günlük saat). Doluluk = gereken
            ÷ mevcut kişi. Çalışma günü = hafta içi − resmî tatil − arife (Cofle'de yarım günde
            çalışılmıyor; kullanıcı 2026-10-21: "W44'te 3 gün çalışacağız"); panelden haftaya elle.
  Süresiz   Ana Veri'de süresi olmayan ürünün yükü bölümün ORTANCA birim süresiyle TAHMİN
            edilir ve ayrı gösterilir (gizlenmez, tanımlıya da karışmaz).

Varsayılan kişi 'Aylık NçS' sayfasından; günlük saat SABİT varsayılan (kullanıcı 2026-10-21:
"kapasiteleri NçS'ye göre değil çalışma gün sayısına göre"). Kişi / günlük saat / verim panelden
bölüm başına (bolum_kapasite_ayar), haftanın gün sayısı haftaya (kapasite_hafta_gun) değiştirilir.
"""
import os
import statistics
from datetime import date, datetime, timedelta

STANDART_GUN = 5                 # bakiye (gecikmiş iş) 'kişi-hafta'sı standart haftayla
VARSAYILAN_SAAT = 8.5
VARSAYILAN_VERIM = 85.0
TK2_SIRA = ('kaynak', 'lazer', 'pres', 'metal', 'montaj')
HARIC_BOLUM = ('TK2|isleme',     # kullanıcı 2026-10-21: "işlemeyi kaldır"
               'TK1|elektronik', 'TK1|bükme')   # 2026-10-22: TK1'de elektronik + bükme iptal
# Ana Veri'deki alt bölüm → kapasitede sayıldığı bölüm (kullanıcı 2026-10-22)
BOLUM_BIRLESTIR = {
    'TK1|joystick': ('TK1|montaj', 'Montaj'),
    'TK1|lf': ('TK1|montaj', 'Montaj'),
    'TK1|grup': ('TK1|montaj', 'Montaj'),
    'TK1|kesim': ('TK1|tel', 'Tel'),
    'TK1|pp': ('TK1|tel', 'Tel'),
}
# Bölüme özgü varsayılanlar (kullanıcı 2026-10-22) — panelden değiştirilir (bolum_kapasite_ayar)
BOLUM_VARSAYILAN = {
    'TK2|lazer': {'birim': 'makine', 'kisi': 4, 'kurulu': 6},   # 6 makinenin 4'ü aynı anda kesimde
    'TK2|kaynak': {'uyum': 50},                                # çift istasyon süre uyumu %50
}
# Resmî tatiller (tam gün) ve arifeler (yarım gün — Cofle'de yarım günde çalışılmıyor, o gün
# çalışma gününden düşer). Dini bayramlar Diyanet takvimine göre girildi; köprü / ek tatil
# ya da farklı bir karar için haftanın gün sayısı panelden elle düzeltilir.
TATILLER = {
    '2026-01-01': 'Yılbaşı',
    '2026-03-19': 'Ramazan Bayramı arifesi', '2026-03-20': 'Ramazan Bayramı',
    '2026-03-21': 'Ramazan Bayramı', '2026-03-22': 'Ramazan Bayramı',
    '2026-04-23': '23 Nisan', '2026-05-01': '1 Mayıs', '2026-05-19': '19 Mayıs',
    '2026-05-26': 'Kurban Bayramı arifesi', '2026-05-27': 'Kurban Bayramı', '2026-05-28': 'Kurban Bayramı',
    '2026-05-29': 'Kurban Bayramı', '2026-05-30': 'Kurban Bayramı',
    '2026-07-15': '15 Temmuz', '2026-08-30': '30 Ağustos',
    '2026-10-28': 'Cumhuriyet Bayramı arifesi', '2026-10-29': 'Cumhuriyet Bayramı',
    '2027-01-01': 'Yılbaşı',
    '2027-03-08': 'Ramazan Bayramı arifesi', '2027-03-09': 'Ramazan Bayramı',
    '2027-03-10': 'Ramazan Bayramı', '2027-03-11': 'Ramazan Bayramı',
    '2027-04-23': '23 Nisan', '2027-05-01': '1 Mayıs',
    '2027-05-15': 'Kurban Bayramı arifesi', '2027-05-16': 'Kurban Bayramı', '2027-05-17': 'Kurban Bayramı',
    '2027-05-18': 'Kurban Bayramı', '2027-05-19': 'Kurban Bayramı · 19 Mayıs',
    '2027-07-15': '15 Temmuz', '2027-08-30': '30 Ağustos',
    '2027-10-28': 'Cumhuriyet Bayramı arifesi', '2027-10-29': 'Cumhuriyet Bayramı',
}
# Bizim üretmediğimiz kaynaklar: kapasiteye girmez (fason / satın alma)
DISARIDA_TESIS = ('satınalma', 'satinalma', 'adp', 'aspanda', 'depo')
HARIC_ONEK = ('92.',)            # 92.* bizde üretilmiyor (feedback 2026-10)

_ONBELLEK = {'anahtar': None, 'veri': None}


def tablolari_kur(conn):
    conn.execute("""
        CREATE TABLE IF NOT EXISTS bolum_kapasite_ayar (
            anahtar TEXT PRIMARY KEY,
            kisi REAL,
            gunluk_saat REAL,
            verim REAL,
            guncelleyen TEXT DEFAULT '',
            updated_at TEXT DEFAULT ''
        )
    """)
    for kolon in ('uyum REAL', 'birim TEXT', 'kurulu REAL'):
        try:
            conn.execute(f'ALTER TABLE bolum_kapasite_ayar ADD COLUMN {kolon}')
        except Exception:
            pass
    conn.execute("""
        CREATE TABLE IF NOT EXISTS kapasite_hafta_gun (
            hafta TEXT PRIMARY KEY,
            gun REAL NOT NULL,
            guncelleyen TEXT DEFAULT '',
            updated_at TEXT DEFAULT ''
        )
    """)


def _tr(s):
    return str(s or '').strip().replace('İ', 'i').replace('I', 'ı').lower()


def _sayi(v):
    try:
        f = float(str(v).replace(',', '.'))
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def _tesis(d):
    s = _tr(d).replace(' ', '').replace('-', '')
    if s in ('tk2', ''):
        return 'TK2' if s else ''
    if s == 'tk1':
        return 'TK1'
    if any(s.startswith(x) for x in DISARIDA_TESIS):
        return None                  # dışarıda
    return str(d).strip()            # Pandora …


def _bolum_ad(tesis, e):
    """(anahtar, görünen ad) — TK2'de Forge bölüm anahtarı, diğerlerinde E'deki ad."""
    import import_excel as IE
    if tesis == 'TK2':
        b = IE.ana_veri_bolum(e)
        return (f'TK2|{b}', IE.ANA_VERI_ETIKET[b]) if b else (None, None)
    ad = ' '.join(str(e or '').split())
    if not ad or any(x in _tr(ad) for x in ('bölüm yok', 'olabilir', 'fason')):
        return None, None
    ad = ad[:1].upper() + ad[1:]
    anahtar = f'{tesis}|{_tr(ad)}'
    return BOLUM_BIRLESTIR.get(anahtar, (anahtar, ad))


def ana_veri_oku(yol=None):
    """Ana Veri → {'urun': {KOD: {...}}, 'ncs': {anahtar: {...}}, 'mtime'} (mtime önbellekli).
    urun: anahtar, tesis, bolum, sn (adet başı), kap1v (vardiya adedi)."""
    import openpyxl
    import import_excel as IE
    yol = yol or IE.ANA_VERI_YOL
    mt = os.path.getmtime(yol)
    if _ONBELLEK['anahtar'] == (yol, mt):
        return _ONBELLEK['veri']
    wb = openpyxl.load_workbook(yol, read_only=True, data_only=True)
    try:
        ws, kol = IE._ana_sayfa_bul(wb)
        if ws is None:
            raise ValueError("Ana Veri sayfası bulunamadı")

        def al(r, a):
            i = kol.get(a)
            return r[i] if (i is not None and i < len(r)) else None
        urun = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            kod = IE._norm_kod(al(r, 'kod'))
            if len(kod) < 2 or kod in urun:
                continue
            tesis = _tesis(al(r, 'tk'))
            if tesis is None:
                urun[kod] = {'disarida': True}
                continue
            anahtar, ad = _bolum_ad(tesis or 'TK2', al(r, 'bolum'))
            if not anahtar:
                urun[kod] = {'bolumsuz': True}
                continue
            g, h = _sayi(al(r, 'cevrim')), _sayi(al(r, 'kaynak'))
            ek = {}
            if anahtar == 'TK2|kaynak' and (g or h):
                sn = (g or 0) + (h or 0)                      # seri: söktak (G) + kaynak (H)
                ek = {'s_sn': g, 'k_sn': h, 'min_sn': min(g, h) if (g and h) else 0.0}
            elif anahtar == 'TK2|metal' and g:
                goz = _sayi(al(r, 'goz')) or 1.0
                sn = g / goz                                  # G baskı başı; bir baskıda göz kadar parça
                ek = {'baski_sn': g, 'goz': goz}
            else:
                sn = g
            urun[kod] = {'anahtar': anahtar, 'tesis': tesis or 'TK2', 'bolum': ad,
                         'sn': sn, 'kap1v': _sayi(al(r, 'kapasite')) if not sn else None, **ek}
        ncs = _ncs_oku(wb)
    finally:
        wb.close()
    veri = {'urun': urun, 'ncs': ncs, 'mtime': datetime.fromtimestamp(mt).strftime('%Y-%m-%d %H:%M')}
    _ONBELLEK.update(anahtar=(yol, mt), veri=veri)
    return veri


def _ncs_oku(wb):
    """'Aylık NçS' → {anahtar: {'kisi', 'ay'}} (varsayılan kişi)."""
    if 'Aylık NçS' not in wb.sheetnames:
        return {}
    ay, out, bas = '', {}, None
    for r in wb['Aylık NçS'].iter_rows(values_only=True):
        r = list(r) + [None] * 6
        if _tr(r[0]) == 'ay':
            ay = str(r[1] or '').strip()[:7]
            continue
        if _tr(r[0]) == 'tesis':
            bas = True
            continue
        if not bas or not r[0] or not r[1]:
            continue
        tesis = _tesis(r[0]) or ''
        e = r[1]
        if tesis == 'TK1' and 'tel' in _tr(e):
            e = 'Tel'
        anahtar, _ad = _bolum_ad(tesis, e)
        kisi = _sayi(r[4])
        if anahtar and kisi:
            out[anahtar] = {'kisi': kisi, 'ay': ay}
    return out


def ayarlar(conn):
    return {r[0]: {'kisi': r[1], 'gunluk_saat': r[2], 'verim': r[3], 'guncelleyen': r[4], 'updated_at': r[5],
                   'uyum': r[6], 'birim': r[7], 'kurulu': r[8]}
            for r in conn.execute("SELECT anahtar, kisi, gunluk_saat, verim, guncelleyen, updated_at, "
                                  "uyum, birim, kurulu FROM bolum_kapasite_ayar")}


def ayar_kaydet(conn, anahtar, kisi=None, gunluk_saat=None, verim=None, kullanici='',
                uyum=None, birim=None, kurulu=None):
    eski = ayarlar(conn).get(anahtar, {})

    def sec(yeni, ad, alt, ust):
        if yeni is None or yeni == '':
            return eski.get(ad)
        v = float(yeni)
        if not alt <= v <= ust:
            raise ValueError(f'{ad} {alt}–{ust} arasında olmalı')
        return v
    if birim not in (None, '') and birim not in ('kisi', 'makine'):
        raise ValueError('birim kisi ya da makine olmalı')
    conn.execute(
        "INSERT INTO bolum_kapasite_ayar (anahtar, kisi, gunluk_saat, verim, guncelleyen, updated_at, uyum, birim, kurulu) "
        "VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(anahtar) DO UPDATE SET kisi=excluded.kisi, "
        "gunluk_saat=excluded.gunluk_saat, verim=excluded.verim, guncelleyen=excluded.guncelleyen, "
        "updated_at=excluded.updated_at, uyum=excluded.uyum, birim=excluded.birim, kurulu=excluded.kurulu",
        (anahtar, sec(kisi, 'kisi', 0, 500), sec(gunluk_saat, 'gunluk_saat', 1, 24),
         sec(verim, 'verim', 10, 150), kullanici, datetime.now().strftime('%Y-%m-%d %H:%M'),
         sec(uyum, 'uyum', 0, 100), birim if birim not in (None, '') else eski.get('birim'),
         sec(kurulu, 'kurulu', 0, 500)))
    conn.commit()


def hafta_gun_kaydet(conn, bas, gun, kullanici=''):
    """Haftanın çalışma günü (bas = pazartesi). gun boş/None → varsayılana dön."""
    date.fromisoformat(bas)                       # geçersiz tarih → ValueError
    if gun is None or str(gun).strip() == '':
        conn.execute("DELETE FROM kapasite_hafta_gun WHERE hafta=?", (bas,))
    else:
        g = float(str(gun).replace(',', '.'))
        if not 0 <= g <= 7:
            raise ValueError('gün 0–7 arasında olmalı')
        conn.execute("INSERT INTO kapasite_hafta_gun (hafta, gun, guncelleyen, updated_at) VALUES (?,?,?,?) "
                     "ON CONFLICT(hafta) DO UPDATE SET gun=excluded.gun, guncelleyen=excluded.guncelleyen, "
                     "updated_at=excluded.updated_at", (bas, g, kullanici, datetime.now().strftime('%Y-%m-%d %H:%M')))
    conn.commit()


def _is_tarihi(bitis):
    """Emrin YÜKÜNÜN yazılacağı gün. Bitişi Pazartesi olan emir bir önceki hafta üretilir
    (kullanıcı 2026-10-21: "pazartesi emirlerini önceki haftaya yaz") → Pazar'a çekilir;
    hafta sonu bitişliler zaten Pzt–Paz haftasının içinde. ValueError: tarih yok."""
    bt = date.fromisoformat(str(bitis)[:10])
    return bt - timedelta(days=1) if bt.weekday() == 0 else bt


def _haftalar(bugun, n, conn=None):
    """[{no, bas, bit, gun, gun_varsayilan, elle, kismi, tatil[]}] — gün: hafta içi − tatil −
    arife, elle girilmişse o. BU HAFTA yalnız BUGÜNDEN itibaren kalan iş günleri sayılır
    (kullanıcı 2026-10-22: "4. günündeyiz, önceki 3 günü dahil etmek mantıklı mı") — geçen
    günlerde yapılan iş teyitle zaten emirden düşmüş; o günlerin kapasitesini saymak işi iki kez
    karşılamak olurdu. Hafta sonundaysak liste gelecek haftadan başlar."""
    pzt = bugun - timedelta(days=bugun.weekday())
    if bugun.weekday() >= 5:
        pzt += timedelta(days=7)
    elle = {}
    if conn is not None:
        try:
            elle = {r[0]: r[1] for r in conn.execute("SELECT hafta, gun FROM kapasite_hafta_gun")}
        except Exception:
            elle = {}
    out = []
    for i in range(n):
        bas = pzt + timedelta(days=7 * i)
        gunler = [bas + timedelta(days=j) for j in range(5)]
        tatil = [f"{g.strftime('%d.%m')} {TATILLER[g.isoformat()]}" for g in gunler if g.isoformat() in TATILLER]
        kismi = bas < bugun                       # hafta başladı: geçen günler sayılmaz
        varsayilan = sum(1 for g in gunler if g.isoformat() not in TATILLER and g >= bugun)
        out.append({'no': f'H{bas.isocalendar()[1]}', 'bas': bas.isoformat(),
                    'bit': (bas + timedelta(days=6)).isoformat(), 'gun_varsayilan': varsayilan,
                    'gun': elle.get(bas.isoformat(), varsayilan), 'elle': bas.isoformat() in elle,
                    'kismi': kismi, 'tatil': tatil})
    return out


def hesapla(conn, hafta=4, bugun=None, yol=None):
    """Bölüm × hafta yük tablosu. Her hücre: kişi-saat (tanımlı + tahmini), adet, kod sayısı.
    Ayar/varsayılan bölüm satırına eklenir; gereken kişi + doluluk İSTEMCİDE hesaplanır
    (kişi kutusu değişince sayfa sunucuya gitmeden güncellensin)."""
    bugun = bugun or date.today()
    av = ana_veri_oku(yol)
    urun = av['urun']
    hft = _haftalar(bugun, hafta, conn)
    sinir = [date.fromisoformat(h['bas']) for h in hft] + [date.fromisoformat(hft[-1]['bit']) + timedelta(days=1)]
    ay = ayarlar(conn)
    bolumler, atanmamis, sayac = {}, {}, {'tarihsiz': 0, 'ufuk_disi': 0, 'disarida': 0}
    talep_ts = None

    def bolum(anahtar, ad, tesis):
        if anahtar not in bolumler:
            a, n, v = ay.get(anahtar, {}), av['ncs'].get(anahtar, {}), BOLUM_VARSAYILAN.get(anahtar, {})
            if a.get('kisi') is not None:
                kisi, kaynak = a['kisi'], 'elle'
            elif v.get('kisi') is not None:
                kisi, kaynak = v['kisi'], 'varsayılan'
            else:
                kisi, kaynak = (n.get('kisi') or 0), ('NçS ' + n['ay'] if n.get('kisi') else '')
            bolumler[anahtar] = {
                'anahtar': anahtar, 'ad': ad, 'tesis': tesis, 'kisi': kisi, 'kisi_kaynagi': kaynak,
                'gunluk_saat': a.get('gunluk_saat') or VARSAYILAN_SAAT,
                'verim': a.get('verim') or VARSAYILAN_VERIM,
                'birim': a.get('birim') or v.get('birim') or 'kisi',
                'kurulu': a.get('kurulu') if a.get('kurulu') is not None else v.get('kurulu'),
                'uyum': a.get('uyum') if a.get('uyum') is not None else v.get('uyum'),
                'guncelleyen': a.get('guncelleyen') or '',
                'hucre': {k: {'saat_sn': 0.0, 'saat_min': 0.0, 'vardiya': 0.0, 'adet': 0.0, 'adet_suresiz': 0.0,
                              'kod': set(), 'suresiz_kod': set()} for k in ['gecikmis'] + [h['no'] for h in hft]},
                '_sn': [], '_kap': [], '_oran': []}
        return bolumler[anahtar]

    for r in conn.execute("SELECT kod, kalan, bitis, senk_at FROM kapasite_talep WHERE kalan > 0"):
        talep_ts = talep_ts or r[3]
        kod = str(r[0] or '').strip().upper().replace(' ', '')
        if not kod or kod.startswith(HARIC_ONEK):
            continue
        u = urun.get(kod)
        if not u or u.get('bolumsuz'):
            a = atanmamis.setdefault(kod, {'kod': kod, 'adet': 0.0, 'emir': 0,
                                           'neden': 'Ana Veri\'de bölümsüz' if u else 'Ana Veri\'de yok'})
            a['adet'] += r[1]
            a['emir'] += 1
            continue
        if u.get('disarida') or u.get('anahtar') in HARIC_BOLUM:
            sayac['disarida'] += 1
            continue
        try:
            bt = _is_tarihi(r[2])
        except ValueError:
            sayac['tarihsiz'] += 1
            continue
        if bt < sinir[0]:
            hk = 'gecikmis'
        elif bt >= sinir[-1]:
            sayac['ufuk_disi'] += 1
            continue
        else:
            hk = next(hft[i]['no'] for i in range(len(hft)) if sinir[i] <= bt < sinir[i + 1])
        b = bolum(u['anahtar'], u['bolum'], u['tesis'])
        h = b['hucre'][hk]
        h['adet'] += r[1]
        h['kod'].add(kod)
        if u.get('sn'):
            h['saat_sn'] += r[1] * u['sn'] / 3600.0
            h['saat_min'] += r[1] * (u.get('min_sn') or 0) / 3600.0     # çift istasyonda örtüşebilen
        elif u.get('kap1v'):
            h['vardiya'] += r[1] / u['kap1v']
        else:
            h['adet_suresiz'] += r[1]
            h['suresiz_kod'].add(kod)

    # Bölümün ortanca birim süresi (süresiz ürünü tahmin etmek için) — Ana Veri'deki
    # TÜM ürünlerinden, yalnız ihtiyacı olanlardan değil
    for u in urun.values():
        b = bolumler.get(u.get('anahtar'))
        if b is not None:
            (b['_sn'] if u.get('sn') else b['_kap'] if u.get('kap1v') else []).append(u.get('sn') or u.get('kap1v'))
            if u.get('sn') and u.get('min_sn'):
                b['_oran'].append(u['min_sn'] / u['sn'])
    for b in bolumler.values():
        b['ortanca_sn'] = round(statistics.median(b['_sn']), 1) if b['_sn'] else None
        b['ortanca_kap1v'] = round(statistics.median(b['_kap']), 1) if b['_kap'] else None
        # Süresiz kaynak ürününün tahmininde de çift istasyon payı: min(K,S) ÷ (K+S) ortancası
        b['ortanca_min_oran'] = round(statistics.median(b['_oran']), 3) if b['_oran'] else 0
        tanimli = suresiz = 0.0
        for h in b['hucre'].values():
            h['kod_sayisi'], h['suresiz_kod_sayisi'] = len(h.pop('kod')), len(h.pop('suresiz_kod'))
            tanimli += h['adet'] - h['adet_suresiz']
            suresiz += h['adet_suresiz']
            h['adet'], h['adet_suresiz'] = round(h['adet']), round(h['adet_suresiz'])
            h['saat_sn'], h['saat_min'], h['vardiya'] = round(h['saat_sn'], 2), round(h['saat_min'], 2), round(h['vardiya'], 3)
        b['kapsama'] = round(100.0 * tanimli / (tanimli + suresiz), 1) if (tanimli + suresiz) else None
        del b['_sn'], b['_kap'], b['_oran']
    sira = {f'TK2|{k}': i for i, k in enumerate(TK2_SIRA)}
    liste = sorted(bolumler.values(), key=lambda b: ({'TK2': 0, 'TK1': 1}.get(b['tesis'], 2),
                                                     sira.get(b['anahtar'], 99), b['ad']))
    # Ana Veri'de olmayan kodlardan yalnız ÜRETİM (P) olanlar uyarıdır — satın alma (A)
    # kodunun OPR'si bizim kapasitemizi ilgilendirmez (kapasite_urun = AS400 BARTF0)
    try:
        prov = {r[0]: r[1] for r in conn.execute("SELECT kod, prov FROM kapasite_urun")}
    except Exception:
        prov = {}
    at = sorted((a for a in atanmamis.values() if prov.get(a['kod'], 'P') == 'P'), key=lambda a: -a['adet'])
    return {'haftalar': hft, 'bolumler': liste, 'ana_veri': av['mtime'], 'talep_ts': talep_ts,
            'varsayilan': {'standart_gun': STANDART_GUN, 'saat': VARSAYILAN_SAAT, 'verim': VARSAYILAN_VERIM},
            'atanmamis': {'kod': len(at), 'adet': round(sum(a['adet'] for a in at)),
                          'ilk': [{**a, 'adet': round(a['adet'])} for a in at[:40]]},
            'sayac': sayac}


def detay(conn, anahtar, hafta_no, hafta=4, bugun=None, yol=None, n=40):
    """Bir bölümün bir haftasındaki ürünler (yüke göre azalan): kod, adet, süre, kişi-saat."""
    bugun = bugun or date.today()
    urun = ana_veri_oku(yol)['urun']
    hft = _haftalar(bugun, hafta)
    bas0 = date.fromisoformat(hft[0]['bas'])
    if hafta_no == 'gecikmis':
        aralik = (date.min, bas0)
    else:
        h = next((x for x in hft if x['no'] == hafta_no), None)
        if not h:
            return []
        aralik = (date.fromisoformat(h['bas']), date.fromisoformat(h['bit']) + timedelta(days=1))
    out = {}
    for r in conn.execute("SELECT kod, kalan, bitis, durum, emir_no FROM kapasite_talep WHERE kalan > 0"):
        kod = str(r[0] or '').strip().upper().replace(' ', '')
        u = urun.get(kod)
        if not u or u.get('anahtar') != anahtar:
            continue
        try:
            bt = _is_tarihi(r[2])
        except ValueError:
            continue
        if not aralik[0] <= bt < aralik[1]:
            continue
        d = out.setdefault(kod, {'kod': kod, 'adet': 0.0, 'emir': 0, 'launch': 0.0, 'en_eski': r[2],
                                 'sn': u.get('sn'), 'kap1v': u.get('kap1v'), 'min_sn': u.get('min_sn') or 0,
                                 'k_sn': u.get('k_sn'), 's_sn': u.get('s_sn'),
                                 'baski_sn': u.get('baski_sn'), 'goz': u.get('goz')})
        d['adet'] += r[1]
        d['emir'] += 1
        if str(r[3]) != '10':
            d['launch'] += r[1]
        d['en_eski'] = min(d['en_eski'], r[2])
    for d in out.values():
        d['saat_sn'] = round(d['adet'] * d['sn'] / 3600.0, 2) if d['sn'] else None
        d['saat_min'] = round(d['adet'] * d['min_sn'] / 3600.0, 2)
        d['vardiya'] = round(d['adet'] / d['kap1v'], 2) if d['kap1v'] else None
        d['adet'], d['launch'] = round(d['adet']), round(d['launch'])
    # Yüke göre: süreli ürünler kişi-saatiyle, süresizler en sonda adetle
    return sorted(out.values(), key=lambda d: (
        -(d['saat_sn'] or (d['vardiya'] or 0) * VARSAYILAN_SAAT), -d['adet']))[:n]
